"""Independent original-input/CSV/Decimal reconciliation and complete Excel readback."""
import subprocess,tempfile,shutil,time
from decimal import Decimal as D
from openpyxl import load_workbook
from q4_utils import *
from microgrid_online import PERIODS,merge_emergency

def close(x,y,tol=1e-6):np.testing.assert_allclose(x,y,rtol=0,atol=tol)
def audit_kind(kind,dates,a,p):
    rows=read_csv(OUT/f'q{kind}_dispatch_10min.csv');assert len(rows)==334*144
    def arr(k):return np.array([float(r[k]) for r in rows]).reshape(334,144)
    for day in range(334):
        part=rows[day*144:(day+1)*144]
        assert [r['date'] for r in part]==[str(dates[day+31])]*144
        assert [int(r['index']) for r in part]==list(range(144)) and [r['period'] for r in part]==PERIODS
    close(arr('load_kw')/6,a[31:,:,0]);close(arr('pv_kw')/6,a[31:,:,1]);close(arr('price_realized'),p[31:],1e-12)
    g,c,d,h,w,e,ee=(arr(k) for k in ('G_effective','C_real','D_real','H_real','W_real','E_start','E_end'))
    balance=g+a[31:,:,1]+d+h-a[31:,:,0]-c-w;state=ee-e-.9*c+d/.9
    close(balance,0);close(state,0);close(ee.ravel()[:-1],e.ravel()[1:],1e-8)
    assert min(e.min(),ee.min())>=1200-1e-7 and max(e.max(),ee.max())<=10800+1e-7
    assert min(x.min() for x in (g,c,d,h,w))>=-1e-7
    assert max(c.max(),d.max())<=5000/6+1e-7 and not np.any((c>1e-7)&(d>1e-7))
    short=a[31:,:,0]-a[31:,:,1]-g
    assert np.all(c<=np.maximum(-short,0)+1e-7) and np.all(d<=np.maximum(short,0)+1e-7)
    close(h,np.maximum(short-d,0));close(w,np.maximum(-short-c,0))
    versions=np.stack([arr(k) for k in ('G0','G6','G12','G18')],axis=1)
    for k in (1,2,3):close(versions[:,k,:k*36],versions[:,k-1,:k*36],0)
    for k in range(4):close(g[:,k*36:(k+1)*36],versions[:,k,k*36:(k+1)*36],0)
    ledger=read_csv(OUT/'q43_adjustment_ledger.csv') if kind==43 else []
    assert len(ledger)==(334*216 if kind==43 else 0)
    if kind==42:close(versions,np.repeat(versions[:,:1],4,axis=1),0)
    commit={(r['date'],int(r['index'])):D(r['G0']) for r in rows};seen=set();up=np.zeros_like(g);down=np.zeros_like(g)
    for r in ledger:
        date=r['delivery_date'];t=int(r['index']);k=int(r['stage']);idx=(date,t);stamp=(date,t,k)
        assert stamp not in seen;seen.add(stamp)
        assert date==r['decision_date'] and t>=k*36 and int(r['decision_hour'])==k*6
        old,new,u,v=(D(r[x]) for x in ('previous_commitment','new_commitment','increase_U','decrease_R'))
        assert abs(commit[idx]-old)<D('1e-6') and abs(new-old-u+v)<D('1e-6')
        assert u>=0 and v>=0 and min(u,v)<D('1e-7');commit[idx]=new
        day=next(j for j,x in enumerate(dates[31:]) if str(x)==date)
        close(float(r['price_realized']),p[day+31,t],1e-12)
        close(float(new),versions[day,k,t]);up[day,t]+=float(u);down[day,t]+=float(v)
        assert abs(D(r['adjustment_charge'])-D(r['price_realized'])*(D('1.5')*u+D('.5')*v))<D('1e-6')
    for r in rows:assert abs(commit[(r['date'],int(r['index']))]-D(r['G_effective']))<D('1e-6')
    costs=dict(plan_cost=p[31:]*arr('G0'),up_cost=1.5*p[31:]*up,penalty=.5*p[31:]*down,refund=np.zeros_like(g),emergency_cost=5*p[31:]*h)
    cash=sum(costs.values());close(arr('cash_yuan'),cash)
    for k,x in costs.items():close(arr(k+'_yuan'),x)
    zero=D(0)
    dec=dict(plan_cost_yuan=sum((D(r['price_realized'])*D(r['G0']) for r in rows),zero),
      emergency_cost_yuan=sum((D(5)*D(r['price_realized'])*D(r['H_real']) for r in rows),zero),
      up_cost_yuan=sum((D('1.5')*D(r['price_realized'])*D(r['increase_U']) for r in ledger),zero),
      penalty_yuan=sum((D('.5')*D(r['price_realized'])*D(r['decrease_R']) for r in ledger),zero))
    dec['total_cost_yuan']=sum(dec.values(),zero)
    summary=json.loads((OUT/f'q{kind}_summary.json').read_text())
    for k,x in dec.items():close(float(x),summary[k])
    daily=read_csv(OUT/f'q{kind}_daily_summary.csv');assert len(daily)==334
    close([float(r['total_cost_yuan']) for r in daily],cash.sum(1))
    merged=merge_emergency(dates[31:],h);saved=read_csv(OUT/f'q{kind}_emergency_purchase_merged.csv');assert len(merged)==len(saved)
    for x,y in zip(merged,saved):assert x['date']==y['date'] and x['period']==y['period'];close(x['kwh'],float(y['kwh']))
    wb=load_workbook(OUT/f'result4-{kind-40}.xlsx',data_only=True,read_only=True)
    template=load_workbook(ROOT/f'附件/附件5/result4-{kind-40}.xlsx',read_only=True)
    assert wb.sheetnames==template.sheetnames
    cells=0
    for name in wb.sheetnames:
        actual=list(wb[name].values);head=list(template[name].values)[0]
        if name in ('计划购电量','调整购电量'):
            assert list(actual[0][1:145])==PERIODS
            assert actual[0][0]==head[0] and actual[0][145:]==head[145:]
        else:assert actual[0]==head,(name,actual[0],head)
        for row in actual:
            assert not any(isinstance(v,str) and v.startswith(('#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#SPILL!')) for v in row)
        if name in ('计划购电量','调整购电量'):
            assert len(actual)==335
            target=arr('G0') if name=='计划购电量' else g
            for j,row in enumerate(actual[1:]):
                assert row[0].date()==dates[j+31]
                close(row[1:145],target[j]);close(row[145],target[j].sum())
                fee=cash[j].sum() if kind==42 or name=='调整购电量' else costs['plan_cost'][j].sum()
                close(row[146],fee);cells+=147
        elif name=='充放电量':
            assert len(actual)==2005
            for j in range(334):
                for b in range(6):
                    row=actual[1+j*6+b]
                    if b==0:assert row[0].date()==dates[j+31]
                    else:assert row[0] is None
                    assert row[1]==f'{b*4}:00-{(b+1)*4}:00'
                    close(row[2],c[j,b*24:(b+1)*24].sum());close(row[3],d[j,b*24:(b+1)*24].sum())
                    if b<2:assert row[4]==('0:00' if b==0 else '24:00');close(row[5],e[j,0] if b==0 else ee[j,-1])
                    else:assert row[4] is None and row[5] is None
                    cells+=6
        else:
            expected=[]
            for date in dates[31:]:
                group=[r for r in merged if r['date']==str(date)]
                expected.extend([(date if i==0 else None,r['period'],r['kwh']) for i,r in enumerate(group)] if group else [(date,'无',0.)])
            assert len(actual)==len(expected)+1
            for row,x in zip(actual[1:],expected):
                assert (row[0].date() if row[0] else None)==x[0] and row[1]==x[1];close(row[2],x[2]);cells+=3
    wb.close();template.close()
    selection=read_csv(OUT/f'experiments/q{kind}_monthly_selection.csv')
    assert len(selection)==26*11 and sum(r['selected']=='True' for r in selection)==11
    cal=json.loads((OUT/f'experiments/q{kind}_calendar.json').read_text())
    from q4_model import Config
    for r in selection:
        day=int(r['selection_day']);assert int(r['end'])==day-1 and int(r['begin'])==day-21
        if r['selected']=='True':
            group=[v for v in selection if v['selection_day']==r['selection_day']]
            close(float(r['score']),min(float(v['score']) for v in group))
            assert Config(**cal[day]).name==r['config']
    formal=dict(np.load(OUT/f'experiments/q{kind}_formal.npz'));close(formal['E'][30,-1],e[0,0])
    for k in ('G','C','D','H','W'):close(formal[k][31:],{'G':g,'C':c,'D':d,'H':h,'W':w}[k])
    return dict(decimal_totals={k:str(v) for k,v in dec.items()},intervals=len(rows),ledger_rows=len(ledger),excel_cells_checked=cells,balance_max=float(abs(balance).max()),state_max=float(abs(state).max()))

def regressions():
    results={}
    for q in (1,2,3,4):
        cmd=[sys.executable,str(ROOT/f'Q{q}/test_q{q}.py')]
        r=subprocess.run(cmd,capture_output=True,text=True);(OUT/f'q{q}_regression.log').write_text(r.stdout+r.stderr)
        results[f'q{q}_tests']=r.returncode;assert r.returncode==0,f'Q{q} tests failed'
    for q in (1,2,3):
        if not (ROOT/f'Q{q}/freeze_q{q}.py').exists():
            results[f'q{q}_freeze']='not_packaged';continue
        r=subprocess.run([sys.executable,str(ROOT/f'Q{q}/freeze_q{q}.py')],capture_output=True,text=True)
        (OUT/f'q{q}_freeze_check.log').write_text(r.stdout+r.stderr);results[f'q{q}_freeze']=r.returncode;assert r.returncode==0
    # Frozen Q3 audit writes its own report: redirect a complete minimal read context to a temporary copy.
    with tempfile.TemporaryDirectory(prefix='q4_old_audit_') as td:
        temp=Path(td)
        for item in ROOT.iterdir():
            if item.name in ('Q3','Q4','.git','.venv','.agents','.codex'):continue
            (temp/item.name).symlink_to(item,target_is_directory=item.is_dir())
        shutil.copytree(ROOT/'Q3',temp/'Q3',ignore=shutil.ignore_patterns('__pycache__','node_modules','.matplotlib','*.zip'))
        r=subprocess.run([sys.executable,'-c',"import sys;sys.path.insert(0,'Q3');from audit_q3 import audit;audit(run_tests=False)"],cwd=temp,capture_output=True,text=True)
        (OUT/'q3_transaction_audit.log').write_text(r.stdout+r.stderr);results['q3_transaction_audit']=r.returncode;assert r.returncode==0
    return results

def audit():
    tic=time.perf_counter();dates,a,fixed,p,records=read_inputs()
    kinds={str(k):audit_kind(k,dates,a,p) for k in (42,43)};tests=regressions();protected=check_protected()
    result=dict(status='PASS',results=kinds,regression_exit_codes=tests,protected_files=protected,seconds=time.perf_counter()-tic)
    write_json(OUT/'validation.json',result)
    (OUT/'Q4最终验收审计.md').write_text('# Q4最终验收审计\n\n独立从原始附件、CSV逐点流量与逐笔调整账本重算，全部通过。summary仅用于核算后交叉检查。\n\n'+json.dumps(result,ensure_ascii=False,indent=2)+f'\n\n两份Excel的所有正式日期、144点计划/有效承诺、每日费用、4小时充放电、日首末SOC及紧急购电全部回读。可用接手保护条目为{protected}；为0时不代表检查了私有原件。not_packaged表示公开衍生包未包含原始冻结校验器，该项未核验。Q3交易审计在临时副本执行，原有审计文件不改写。\n\n限制：未来价格未知、当期聚合可见、交付价结算为建模解释；E1的阶段VOI采用E0代理；334日单市场测试不构成普适统计优势。哈希及数值验算不是外部独立认证。本次不包含私有冻结件完整性证明或人工视觉核对。\n')
    print('Q4 independent audit PASS',flush=True)
if __name__=='__main__':audit()

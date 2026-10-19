"""Public Python exporter for the verified Q4 payload; no numerical policy changes."""
from pathlib import Path
from datetime import datetime
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
OUT=Path(__file__).resolve().parent

def export():
    for kind in (42,43):
        name=f'result4-{kind-40}.xlsx'
        wb=load_workbook(OUT.parent/'附件/附件5'/name)
        payload=json.loads((OUT/f'q{kind}_workbook_payload.json').read_text())
        for title,rows in payload.items():
            ws=wb[title]
            for merged in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged))
            for row in ws:
                for cell in row:cell.value=None
            for i,row in enumerate(rows,1):
                for j,value in enumerate(row,1):
                    if i>1 and j==1 and isinstance(value,str) and len(value)==10:
                        value=datetime.strptime(value,'%Y-%m-%d')
                    cell=ws.cell(i,j,value)
                    cell.font=Font(name='Arial',size=11)
                    if i>1:cell.number_format='yyyy-mm-dd' if j==1 else '#,##0.0000'
                    if i==1:cell.alignment=Alignment(wrap_text=True)
            ws.row_dimensions[1].height=34
            ws.column_dimensions['A'].width=15
            for j in range(2,len(rows[0])+1):ws.column_dimensions[get_column_letter(j)].width=17 if len(rows[0])>10 else 24
            ws.freeze_panes='A2'
            if title=='充放电量':
                for d in range(334):ws.merge_cells(start_row=2+d*6,end_row=7+d*6,start_column=1,end_column=1)
        wb.properties.creator='CUMCM Team';wb.properties.lastModifiedBy='CUMCM Team'
        wb.save(OUT/name)
if __name__=='__main__':export()

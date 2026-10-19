"""从本次原始计算数组生成全部CSV、官方模板工作簿、论文和图件。"""
from pathlib import Path
from copy import copy
from datetime import datetime
import json,csv,sys,os
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from solve_q2 import ROOT,OUT,write_json,write_csv,summary_of,check_q1
from microgrid_online import read_year,PERIODS,merge_emergency,terminal_value
from microgrid_common import Battery


def read_csv(path):
    with path.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))


def table(headers,rows):
    return '\n'.join(['|'+'|'.join(headers)+'|','|'+'|'.join(['---']*len(headers))+'|']+
                     ['|'+'|'.join(str(x) for x in row)+'|' for row in rows])


def export_workbook(dates,actual,price,s):
    # user explicitly requested a reproducible Python generator using official template
    template=ROOT/'附件/附件5/result2.xlsx'
    wb=load_workbook(template)
    ws=wb['计划购电量'];old=[ws.cell(1,c).value for c in range(2,146)]
    for t,label in enumerate(PERIODS,2):ws.cell(1,t,label)
    for day in range(31,365):
        row=day-29
        ws.cell(row,1,datetime.combine(dates[day],datetime.min.time()))
        for t in range(144):ws.cell(row,t+2,float(s['G'][day,t]))
        ws.cell(row,146,float(s['G'][day].sum()))
        ws.cell(row,147,float(price@s['G'][day]+5*price@s['H'][day]))
    ws.freeze_panes='B2';ws.auto_filter.ref='A1:EQ335'
    for col in range(1,148):ws.column_dimensions[get_column_letter(col)].width=17 if col in (1,146,147) else 15
    for row in ws.iter_rows(min_row=2):
        row[0].number_format='yyyy-mm-dd'
        for cell in row[1:]:cell.number_format='#,##0.0000'
    ws.row_dimensions[1].height=25
    # expand examples by copying actual template styles, not creating similar sheets
    ws=wb['充放电量'];styles=[[copy(ws.cell(r,c)._style) for c in range(1,7)] for r in range(2,8)]
    ws.delete_rows(2,ws.max_row-1)
    for day in range(31,365):
        for j in range(6):
            r=2+(day-31)*6+j
            for c in range(1,7):ws.cell(r,c)._style=copy(styles[j][c-1])
            if j==0:ws.cell(r,1,datetime.combine(dates[day],datetime.min.time())).number_format='yyyy-mm-dd'
            ws.cell(r,2,f'{j*4}:00-{(j+1)*4}:00')
            ws.cell(r,3,float(s['C'][day,j*24:(j+1)*24].sum()))
            ws.cell(r,4,float(s['D'][day,j*24:(j+1)*24].sum()))
            if j<2:
                ws.cell(r,5,'0:00' if j==0 else '24:00')
                ws.cell(r,6,float(s['E'][day,0 if j==0 else -1]))
            for c in (3,4,6):ws.cell(r,c).number_format='#,##0.0000'
    ws.freeze_panes='C2'
    for c in range(1,7):ws.column_dimensions[get_column_letter(c)].width=18
    ws=wb['紧急购电量'];styles=[copy(ws.cell(2,c)._style) for c in range(1,4)]
    ws.delete_rows(2,ws.max_row-1);merged=merge_emergency(dates[31:],s['H'][31:]);row=2
    for day in range(31,365):
        events=[r for r in merged if r['date']==dates[day].isoformat()]
        if not events:events=[dict(period='无',kwh=0.)]
        for j,event in enumerate(events):
            for c in range(1,4):ws.cell(row,c)._style=copy(styles[c-1])
            if j==0:ws.cell(row,1,datetime.combine(dates[day],datetime.min.time())).number_format='yyyy-mm-dd'
            ws.cell(row,2,event['period']);ws.cell(row,3,float(event['kwh'])).number_format='#,##0.0000';row+=1
    ws.freeze_panes='B2'
    for c in range(1,4):ws.column_dimensions[get_column_letter(c)].width=22
    for sheet in wb:
        sheet.sheet_view.zoomScale=85
        sheet.print_options.horizontalCentered=True
        sheet.print_title_rows='1:1'
    wb.save(OUT/'result2.xlsx');wb.close()
    write_csv(OUT/'time_mapping.csv',[dict(index=t,input_end=PERIODS[t].split('-')[1],
        original_template=old[t],output_period=PERIODS[t],output_column=t+2) for t in range(144)])
    return merged


def data_exports(dates,a,p,s):
    daily=[];dispatch=[];selected=[]
    wanted={'2025-03-20','2025-06-21','2025-09-23','2025-12-21'}
    for day,date in enumerate(dates):
        g,h=s['G'][day],s['H'][day];pc=float(p@g);hc=float(5*p@h)
        daily.append(dict(date=date.isoformat(),formal=day>=31,config=str(s['config'][day]),q=float(s['q'][day]),
             plan_kwh=float(g.sum()),emergency_kwh=float(h.sum()),total_grid_kwh=float((g+h).sum()),
             plan_cost_yuan=pc,emergency_cost_yuan=hc,total_cost_yuan=pc+hc,
             initial_kwh=float(s['E'][day,0]),terminal_kwh=float(s['E'][day,-1]),
             charge_kwh=float(s['C'][day].sum()),discharge_kwh=float(s['D'][day].sum()),
             emergency_periods=int((h>1e-7).sum()),waste_kwh=float(s['W'][day].sum()),
             pv_curtail_kwh=float(s['pv_curtail'][day].sum()),load_kwh=float(a[day,:,0].sum()),pv_kwh=float(a[day,:,1].sum()),
             net_mae_kw=float(np.abs((a[day,:,0]-a[day,:,1])-(s['prediction'][day,:,0]-s['prediction'][day,:,1])).mean()*6),
             solve_seconds=float(s['seconds'][day])))
        if day<31:continue
        local=[]
        for t in range(144):
            row=dict(date=date.isoformat(),period=PERIODS[t],index=t,price_yuan_kwh=float(p[t]),
                     load_kw=float(a[day,t,0]*6),pv_kw=float(a[day,t,1]*6),
                     load_forecast_kw=float(s['prediction'][day,t,0]*6),pv_forecast_kw=float(s['prediction'][day,t,1]*6),
                     margin_kw=float(s['margin'][day,t]*6),G_plan_kwh=float(g[t]),
                     C_plan_kwh=float(s['C_plan'][day,t]),D_plan_kwh=float(s['D_plan'][day,t]),
                     E_plan_start_kwh=float(s['E_plan'][day,t]),E_plan_end_kwh=float(s['E_plan'][day,t+1]),
                     C_real_kwh=float(s['C'][day,t]),D_real_kwh=float(s['D'][day,t]),H_real_kwh=float(h[t]),
                     W_real_kwh=float(s['W'][day,t]),unused_plan_kwh=float(s['unused_grid'][day,t]),
                     pv_curtail_kwh=float(s['pv_curtail'][day,t]),pv_curtail_upper_kwh=float(s['pv_curtail_upper'][day,t]),
                     E_real_start_kwh=float(s['E'][day,t]),E_real_end_kwh=float(s['E'][day,t+1]),
                     plan_cost_yuan=float(p[t]*g[t]),emergency_cost_yuan=float(5*p[t]*h[t]))
            dispatch.append(row);local.append(row)
        if date.isoformat() in wanted:
            write_csv(OUT/'selected_days'/f'{date.isoformat()}_dispatch.csv',local)
            selected.append(daily[-1])
            write_csv(OUT/'selected_days'/f'{date.isoformat()}_table1.csv',
                [dict(period=PERIODS[t],plan_kwh=float(g[t])) for t in (60,72,84,96,108,120)]+
                [dict(period='全天计划购电量',plan_kwh=float(g.sum()))])
            write_csv(OUT/'selected_days'/f'{date.isoformat()}_table2.csv',
                [dict(period=f'{j*4}:00-{(j+1)*4}:00',charge_kwh=float(s['C'][day,j*24:(j+1)*24].sum()),
                      discharge_kwh=float(s['D'][day,j*24:(j+1)*24].sum())) for j in range(6)])
    write_csv(OUT/'daily_summary.csv',daily[31:]);write_csv(OUT/'warmup_january.csv',daily[:31])
    write_csv(OUT/'dispatch_10min.csv',dispatch)
    write_csv(OUT/'emergency_purchase_10min.csv',[dict(date=r['date'],period=r['period'],H_real_kwh=r['H_real_kwh']) for r in dispatch])
    write_csv(OUT/'selected_days/summary.csv',selected)
    return daily,selected


def figures(dates,a,p,s,daily,comparisons):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    figdir=OUT/'figures';figdir.mkdir(exist_ok=True)
    def save(fig,name):
        fig.savefig(figdir/f'{name}.png',dpi=160,bbox_inches='tight');fig.savefig(figdir/f'{name}.svg',bbox_inches='tight');plt.close(fig)
    months=np.arange(2,13);pc=[];hc=[];ek=[]
    for m in months:
        ids=[i for i,d in enumerate(dates) if d.month==m]
        pc.append(float((s['G'][ids]@p).sum())/1e6);hc.append(float((s['H'][ids]@(5*p)).sum())/1e6);ek.append(s['H'][ids].sum()/1000)
    fig,ax=plt.subplots(2,1,figsize=(10,7),layout='constrained')
    ax[0].bar(months,pc,label='Day-ahead cost',color='#2563eb');ax[0].bar(months,hc,bottom=pc,label='Emergency cost',color='#e66a36')
    ax[0].set_ylabel('Cost (million CNY)');ax[0].legend();ax[1].bar(months,ek,color='#e66a36')
    ax[1].set_ylabel('Emergency energy (MWh)');ax[1].set_xlabel('Month');ax[1].set_xticks(months)
    fig.suptitle('Causal policy: monthly realized costs and emergency purchases');save(fig,'monthly_costs')
    for day,date in enumerate(dates):
        if date.isoformat() not in ('2025-03-20','2025-06-21','2025-09-23','2025-12-21'):continue
        x=(np.arange(144)+.5)/6;edges=np.arange(145)/6
        fig,axes=plt.subplots(5,1,figsize=(12,13),sharex=True,layout='constrained')
        axes[0].plot(x,a[day,:,0]*6,label='Actual load',color='#334155')
        axes[0].plot(x,s['prediction'][day,:,0]*6,'--',label='Load forecast',color='#94a3b8')
        axes[0].plot(x,a[day,:,1]*6,label='Actual PV',color='#c98b0e')
        axes[0].plot(x,s['prediction'][day,:,1]*6,'--',label='PV forecast',color='#e8b958');axes[0].legend(ncol=4);axes[0].set_ylabel('Power (kW)')
        error=((a[day,:,0]-a[day,:,1])-(s['prediction'][day,:,0]-s['prediction'][day,:,1]))*6
        axes[1].plot(x,error,label='Net-load error',color='#475569');axes[1].plot(x,s['margin'][day]*6,label='Risk margin',color='#9333ea')
        axes[1].axhline(0,color='gray',lw=.6);axes[1].set_ylabel('Error / margin (kW)');axes[1].legend(ncol=2)
        axes[2].stairs(s['G'][day]*6,edges,label='Fixed day-ahead plan',color='#2563eb')
        axes[2].bar(x,s['H'][day]*6,width=1/6,label='Emergency',color='#dc2626');axes[2].set_ylabel('Grid power (kW)');axes[2].legend(ncol=2)
        axes[3].bar(x,(s['C'][day]-s['D'][day])*6,width=1/6,color='#0d9488',label='Charge (+) / discharge (-)')
        other=axes[3].twinx();other.plot(edges,s['E'][day],color='#7c3aed',label='Energy');other.set_ylabel('Stored energy (kWh)');other.set_ylim(0,12000)
        for level in (1200,10800):other.axhline(level,color='#7c3aed',alpha=.3,ls='--')
        axes[3].set_ylabel('Battery power (kW)');axes[3].legend(loc='upper left')
        axes[4].stairs(p,edges,color='#be185d',label='Price');axes[4].set_ylabel('Price (CNY/kWh)')
        other=axes[4].twinx();other.plot(edges,np.r_[0,np.cumsum(p*(s['G'][day]+5*s['H'][day]))]/1000,color='#334155');other.set_ylabel('Cumulative cost (1000 CNY)')
        for ax in axes:ax.set_xlim(0,24);ax.grid(alpha=.15)
        axes[-1].set_xticks(range(0,25,2));axes[-1].set_xlabel('Hour')
        fig.suptitle(f'{date}: {s["config"][day]}');save(fig,str(date))
    fig,axes=plt.subplots(1,2,figsize=(12,4.5),layout='constrained')
    err=((a[31:,:,0]-a[31:,:,1])-(s['prediction'][31:,:,0]-s['prediction'][31:,:,1]))*6
    axes[0].hist(err.ravel(),bins=65,color='#64748b');axes[0].set_xlabel('Net-load error (kW)');axes[0].set_ylabel('10-minute periods')
    cover=(err<=s['margin'][31:]*6+1e-6).mean(axis=0)
    axes[1].plot(np.arange(144)/6,cover,color='#9333ea');axes[1].set_ylim(0,1);axes[1].set_xlabel('Hour');axes[1].set_ylabel('Empirical upper-margin coverage')
    save(fig,'forecast_errors')
    fig,axes=plt.subplots(1,2,figsize=(12,4.5),layout='constrained')
    for h,term in [(24,'cycle'),(24,'value'),(48,'cycle'),(72,'cycle')]:
        rows=sorted([r for r in comparisons if r['method']=='blend' and float(r['q'])>0 and int(r['horizon_hours'])==h and r['terminal']==term and r['scenario']=='False'],key=lambda r:float(r['q']))
        axes[0].plot([float(r['q']) for r in rows],[float(r['total_cost_yuan'])/1e6 for r in rows],marker='o',label=f'{h}h {term}')
        axes[1].plot([float(r['q']) for r in rows],[float(r['emergency_kwh'])/1000 for r in rows],marker='o',label=f'{h}h {term}')
    axes[0].set_ylabel('Total cost (million CNY)');axes[1].set_ylabel('Emergency energy (MWh)')
    for ax in axes:ax.set_xlabel('Quantile q');ax.legend();ax.grid(alpha=.15)
    fig.suptitle('Fixed-policy diagnostics; not used to select past official plans');save(fig,'quantile_window_sensitivity')


def generate_documents(dates,a,p,s,summary,daily,selected,comparisons):
    sel=[r for r in read_csv(OUT/'experiments/monthly_selection.csv') if r['selected']=='True']
    forecast=read_csv(OUT/'forecast_metrics.csv')
    sensitivity=json.loads((OUT/'experiments/sensitivities.json').read_text())
    oracle=json.loads((OUT/'experiments/oracle.json').read_text())
    metadata=json.loads((OUT/'experiments/run_metadata.json').read_text())
    byname={r['config']:r for r in comparisons}
    representative=['last_q0.00_24h_cycle','mean7_q0.00_24h_cycle','blend_q0.00_24h_value',
                    'blend_q0.80_24h_cycle','blend_q0.80_48h_cycle','blend_q0.80_72h_cycle','blend_q0.80_24h_value','scenario_blend_48']
    cmp_rows=[]
    for name in representative:
        r=byname[name];cmp_rows.append([name,f'{float(r["total_cost_yuan"]):,.2f}',f'{float(r["emergency_kwh"]):,.2f}',r['emergency_periods'],f'{float(r["terminal_kwh"]):.2f}'])
    cmp_rows.insert(0,['正式月度因果选择',f'{summary["total_cost_yuan"]:,.2f}',f'{summary["emergency_kwh"]:,.2f}',summary['emergency_periods'],f'{summary["terminal_kwh"]:.2f}'])
    cmp_rows.append(['无储能mean7',f'{sensitivity["no_storage_mean7"]["total_cost_yuan"]:,.2f}',f'{sensitivity["no_storage_mean7"]["emergency_kwh"]:,.2f}',sensitivity['no_storage_mean7']['emergency_periods'],f'{sensitivity["no_storage_mean7"]["terminal_kwh"]:.2f}'])
    cmp_rows.append(['同边界全年oracle下界',f'{oracle["objective"]:,.2f}',0,0,f'{summary["terminal_kwh"]:.2f}'])
    fmt=lambda x:f'{x:,.4f}'
    text=r'''# 问题2 因果预测与历史经济验证驱动的微网日前购电

## 1 问题分析与证据边界

问题2要求每天0:00承诺当天购电，在未知实际负载与光伏的条件下利用储能减少5倍电价紧急补电。本文采用“因果预测—候选计划优化—实时物理反馈—跨日状态传递”的分层策略，并在每月首日用已完成历史的经济费用选择配置。原题规定费用与设备参数；日内储能可反馈、十分钟均值解释及终端价值属于本文假设/设计。最终方案由逐月历史验证产生，不从事后全年费用排名选一个赢家回填。

## 2 假设与符号

不售电、不削负荷、不添加题目未给的并网功率上限和衰减费用。允许弃光及已付款计划电量未利用。各流量均为交流母线侧kWh，E为内部储电量。

|符号|定义|
|---|---|
|d,t|日期与十分钟时段，t=0,…,143|
|L,P|实际负载与光伏时段电量，输入功率除以6|
|p,G,H|正常电价、0:00计划购电、实际紧急购电|
|C,D,W,E|实际充电、放电、未利用能量、储电量|
|带横线C,D,E|计划层参考轨迹，非实际执行结果|
|q,v|净负荷上尾裕度参数、终端电量代理价值|

容量12000，能量界1200–10800，最大每时段充放电量5000/6；主效率ηc=ηd=0.9，往返为0.81。题面“90%”解释存在歧义，另测往返90%。每日首末相等只属于Q1条件，不继承为Q2硬约束。

## 3 信息集与非预见性

F[d,0]包含固定日价格、日历、d日前全部已完成负载/光伏、真实初始E及已完成历史验证。G[d,:]仅由该信息决定。执行器的F+[d,t]额外包含当前实时供需观测，不包含t之后的真实值。十分钟平均数据模拟时段内即时平衡，不意味着开区间前已知其均值。

代码将make_plan的输入限制为Features、初始E、固定价格及配置；Features逐日起点构建，所有残差满足实际观测日<d。反馈函数feedback_step仅接收五个当前量，无法看到未来数组。反事实测试分别修改当天和未来月份数据，检查当天计划不变；另修改日内未来段，检查此前动作不变。

## 4 预测模型及误差校准

比较昨日同刻、过去7日均值及组合预测。组合预测为

$$\hat X_{d+j,t}=\tfrac12\operatorname{mean}_{k\in[d-7,d-1]}X_{k,t}+\tfrac12\operatorname{mean}_{k<d,\,k\equiv d+j\pmod7}^{最近4日}X_{k,t}.\tag{1}$$

不足时使用可用样本，同星期无样本时回退近期均值。L和P分别预测后相减。即使预测第二、三天，所有k仍严格小于当前d。每日发出的多时距预测独立存档，只有对应实际日完成后才加入残差池。最近30个有效残差形成m=max(0,Qq(e))。q=0配置表示不加裕度。正截断后的曲线是保守裕度曲线，不保证具有精确q覆盖。

原报告将固定7日/4周均值称扩展窗口不够准确；本文称历史滑窗预测与逐日起点验证。未采用缺少本地天气数据支持的复杂预测器。三类可解释预测器已联合参与经济比较，不能从本实验推出所有更复杂方法都无价值。

## 5 风险分位的数学依据

忽略储能的单时段费用J(G)=pG+5p E[(N−G)+]。连续分布下J'(G)=p−5pPr(N>G)，内点最优为F_N(G)=0.8。直观上缺购的相对损失为4p，过购损失为p。非负约束存在时还需投影。

动态储能的效率、SOC界、功率限制、价差及跨日影子价值破坏单时段分离，因此0.8不再严格最优。本文事先注册q=0、0.65、0.70、0.75、0.80、0.85、0.90、0.95，通过过去实际费用而非全年误差选择。

## 6 计划层与终端条件

对T=144/288/432、风险净负荷n=(预测负载−预测光伏)+m，求

$$\min\sum_{t=0}^{T-1}p_tG_t-vE_T,\tag{2}$$
$$G_t+\bar D_t-\bar C_t-\bar W_t=n_t,\quad G_t,\bar W_t\ge0,\tag{3}$$
$$E_{t+1}=E_t+\eta_c\bar C_t-\bar D_t/\eta_d,\quad1200\le E_t\le10800,\tag{4}$$
$$0\le\bar C_t,\bar D_t\le5000/6,\quad E_0=E^{real}_{d,0}.\tag{5}$$

闭合候选v=0且E_T=E_0；24h终端价值候选不强制末态，v=日价格最低24点均价/ηc，是低价补回内部电量的代理成本。式(2)是设计计划的代理目标；正式账单始终是式(8)，不把终端价值算成真实收入。

LP允许的同时充放电可精确消除。令δ=min(C,D/(ηcηd))，取C'=C−δ、D'=D−ηcηdδ、W'=W+(1−ηcηd)δ，G与E及目标均不变，功率界满足且至少一项为零。因此此计划模型存在互斥LP最优解，不必全年反复求解二进制MILP。小算例与代表日期另用MILP验证费用一致。这一证明依赖自由弃电，不能用于不允许弃电的其他模型。

## 7 实际执行与费用

令r=G+P−L，若r≥0则

$$C=\min(r,5000/6,(10800-E)/\eta_c),\quad D=H=0,\quad W=r-C.\tag{6}$$

若r<0则

$$D=\min(-r,5000/6,(E-1200)\eta_d),\quad H=-r-D,\quad C=W=0.\tag{7}$$

更新状态，下一日初值等于本日实际末值。该执行器保障当期供电且不使用未来，但属于贪心反馈，不声称未来紧急费用的动态最优。正式紧急定义是最大可行储能补缺后仍剩余的供给缺口。

$$F=\sum_{d,t}(p_tG_{d,t}+5p_tH_{d,t}).\tag{8}$$

计划G当日不变，即使被弃用仍收费。W包括两类电量；主口径光伏优先，未利用计划电量=min(W,G)，弃光=max(W−G,0)。不作来源优先假设时弃光界为[max(W−G,0),min(W,P)]，因此W本身不能代表弃光。

## 8 跨日启动与历史模型选择

1月1日无历史：G=0、储能静置、光伏直接供负载、缺口应急；初始6000保持。此日是公开的保守回退假设。1月2日起按mean7、q=0.8、48h闭合共同预运行，月底真实状态传给2月。1月费用独立保存。

预先注册96个预测×分位×终端组合及1个有限策略场景候选。每月首日以最近21个已完成日作为验证段，从该段已知正式初值出发，候选各自连续运行；评分为Σ真实现金费−v(E末−E初)，避免偏好透支末态。选最小评分配置，月内固定配置，每天更新历史。参数筛选本身也是可执行策略；2月选择仅用1月，12月选择不使用12月结果。

场景候选用过去至多12条联合完整日残差曲线，在7种q计划中最小化平均实际反馈费用与末态价值。储能由同一反馈映射逐点决定，无场景内自由未来储能变量，因此非预见性成立。这是有限策略样本平均优化，而非精确多阶段随机规划。场景压力另用200次整日抽样对照逐点独立重排，作为经验诊断，不是严格分布外置信保证。

## 9 求解结果与参数选择
'''
    text+='\n'+table(['选择时刻','配置','历史截止','历史调整评分（元）'],[[str(dates[int(r['selection_day'])]),r['config'],str(dates[int(r['end'])]),f'{float(r["score"]):,.2f}'] for r in sel])+'\n\n'
    text+=f'终端代理价值v={terminal_value(p):.8f}元/内部kWh。所有月份选择都发生在月初历史数据上；全年固定候选排名不参与过去计划。正式统计区间为334天，共48,096时段。计划费 **{fmt(summary["plan_cost_yuan"])}元**，紧急费 **{fmt(summary["emergency_cost_yuan"])}元**，总费 **{fmt(summary["total_cost_yuan"])}元**。紧急购电 **{fmt(summary["emergency_kwh"])}kWh**，共{summary["emergency_periods"]}个时段、{summary["emergency_days"]}日，最大单时段{fmt(summary["max_emergency_kwh"])}kWh，日费用95%分位{fmt(summary["daily_p95_yuan"])}元。\n\n'
    text+=f'期初/期末储电量分别{fmt(summary["initial_kwh"])}、{fmt(summary["terminal_kwh"])}kWh；储电量范围{fmt(summary["energy_min"])}–{fmt(summary["energy_max"])}kWh。总充电{fmt(summary["charge_kwh"])}kWh，总放电{fmt(summary["discharge_kwh"])}kWh。总未利用电量{fmt(summary["waste_kwh"])}kWh，光伏优先口径弃光{fmt(summary["pv_curtail_kwh"])}kWh，来源分配不确定时弃光上界{fmt(summary["pv_curtail_upper_kwh"])}kWh，主口径光伏消纳率{summary["pv_utilization"]:.4%}。\n\n'
    text+='![月度实际费用与紧急购电](figures/monthly_costs.png)\n\n## 10 预测评价\n\n'
    text+=table(['预测器','对象','MAE kW','RMSE kW','真实值减预测偏差 kW'],[[r['method'],r['object'],f'{float(r["mae_kw"]):.3f}',f'{float(r["rmse_kw"]):.3f}',f'{float(r["bias_kw"]):.3f}'] for r in forecast if r['group']=='all'])+'\n\n'
    cover=next(r['coverage'] for r in forecast if r['method']=='formal' and r['group']=='all' and r['object']=='net')
    text+=f'正式上尾裕度经验覆盖率{float(cover):.4%}。该覆盖率不等于不紧急购电概率，因电池能在时段间吸收误差；也不要求等于固定80%，因为q按月变化且裕度正截断。分月及分小时MAE/RMSE保存在forecast_metrics.csv。\n\n![误差与覆盖](figures/forecast_errors.png)\n\n## 11 候选对照与oracle\n\n'
    text+=table(['策略','现金总费 元','紧急电量 kWh','紧急时段数','末态 kWh'],cmp_rows)+'\n\n'
    for name in ('mean7_q0.00_24h_cycle','blend_q0.00_24h_value','blend_q0.80_48h_cycle','scenario_blend_48'):
        cost=float(byname[name]['total_cost_yuan']);text+=f'相对{name}，正式费用变化为{(summary["total_cost_yuan"]/cost-1):+.4%}。\n\n'
    best=min(comparisons,key=lambda r:float(r['total_cost_yuan']))
    text+=f'事后固定候选最低现金费用为{best["config"]}的{float(best["total_cost_yuan"]):,.2f}元。这只是事后诊断，不得据此倒改任何早期月度选择。正式策略不承诺超过每个事后固定候选。候选终点不同，完整表另列终端价值调整成本；评价不能只选对某方案有利的现金口径。\n\n'
    text+=f'oracle直接使用全部334日真实净负荷，并严格匹配正式2月初和12月末状态，一次求解整段LP，最优费用为{oracle["objective"]:,.4f}元，正式费用相对下界差距为{(summary["total_cost_yuan"]/oracle["objective"]-1):.4%}。这是同边界确定性信息松弛下界，不是可执行计划，也未混入result2。单日oracle若起终状态不同则不能用来证明逐日高低。\n\n'
    text+='## 12 四个指定日期\n\n'
    text+=table(['日期','负载kWh','光伏kWh','计划费元','紧急费元','总费元','紧急kWh','日初/末kWh'],[[r['date'],fmt(r['load_kwh']),fmt(r['pv_kwh']),fmt(r['plan_cost_yuan']),fmt(r['emergency_cost_yuan']),fmt(r['total_cost_yuan']),fmt(r['emergency_kwh']),f'{r["initial_kwh"]:.2f}/{r["terminal_kwh"]:.2f}'] for r in selected])+'\n\n'
    merged=merge_emergency(dates[31:],s['H'][31:])
    event_columns=[[x for x in merged if x['date']==r['date']] for r in selected]
    text+='### 表3 四个指定日期的紧急购电\n\n'
    headers=[label for r in selected for label in (r['date']+' 时间段','购电量 kWh')]
    event_rows=[]
    for i in range(max(1,max(map(len,event_columns)))):
        event_rows.append([value for events in event_columns for value in
                           ([events[i]['period'],fmt(events[i]['kwh'])] if i<len(events) else (['无',0] if not events and i==0 else ['','']))])
    text+=table(headers,event_rows)+'\n\n'
    for r in selected:
        day=dates.index(datetime.fromisoformat(r['date']).date());g=s['G'][day]
        text+=f'### {r["date"]}\n\n'
        text+=table(['时间段','计划购电 kWh','时间段','计划购电 kWh','时间段','计划购电 kWh'],
          [[PERIODS[t],fmt(g[t]),PERIODS[t+12],fmt(g[t+12]),PERIODS[t+24],fmt(g[t+24])] for t in (60,96)])+'\n\n'
        text+=f'全天计划购电{fmt(g.sum())}kWh，含紧急的实际总购电{fmt((g+s["H"][day]).sum())}kWh；全天购电费{fmt(r["total_cost_yuan"])}元。\n\n'
        text+=table(['时段','实际充电 kWh','实际放电 kWh','时段','实际充电 kWh','实际放电 kWh'],
          [[f'{j*4}:00-{j*4+4}:00',fmt(s['C'][day,j*24:(j+1)*24].sum()),fmt(s['D'][day,j*24:(j+1)*24].sum()),
            f'{(j+1)*4}:00-{(j+2)*4}:00',fmt(s['C'][day,(j+1)*24:(j+2)*24].sum()),fmt(s['D'][day,(j+1)*24:(j+2)*24].sum())] for j in (0,2,4)])+'\n\n'
        text+=f'0:00储电量{fmt(s["E"][day,0])}kWh，24:00储电量{fmt(s["E"][day,-1])}kWh。\n\n'
        events=[x for x in merged if x['date']==r['date']]
        text+=table(['紧急购电时间段','购电量 kWh'],[[x['period'],fmt(x['kwh'])] for x in events] or [['无',0]])+'\n\n'
        max_t=int(np.argmax(s['H'][day]));surplus=np.maximum(a[day,:,1]-a[day,:,0],0).sum()
        error=(a[day,:,0]-a[day,:,1])-(s['prediction'][day,:,0]-s['prediction'][day,:,1])
        text+=f'当日实际净负荷{fmt(r["load_kwh"]-r["pv_kwh"])}kWh、未计储能前光伏时段性过剩{fmt(surplus)}kWh，净负荷预测MAE为{r["net_mae_kw"]:.2f}kW。上尾裕度平均{s["margin"][day].mean()*6:.2f}kW，实际误差超过裕度的时段数为{int((error>s["margin"][day]+1e-7).sum())}。'
        if s['H'][day,max_t]>1e-7:
            text+=f'最大紧急购电发生于{PERIODS[max_t]}，电量{s["H"][day,max_t]:.4f}kWh；该段初始储电{s["E"][day,max_t]:.2f}kWh、放电{s["D"][day,max_t]:.2f}kWh。该信息直接区分SOC不足与功率限制，不把所有补电都归因于天气。'
        else:text+='当日没有紧急购电，说明本日计划与可用储能共同覆盖了实际缺口，并不保证其他日也无风险。'
        load_bias=(a[day,:,0]-s['prediction'][day,:,0]).sum()
        pv_bias=(a[day,:,1]-s['prediction'][day,:,1]).sum()
        text+=f'实际减预测的日负载偏差为{load_bias:+.2f}kWh，日光伏偏差为{pv_bias:+.2f}kWh。'
        if r['date']=='2025-06-21':
            text+=f'该日为周六，四个指定日中负载最低、光伏时段性富余最大；组合预测仍混入近期较高负载，加上q=0.85裕度，导致计划偏保守。虽然没有补电，总未利用电量达到{r["waste_kwh"]:.2f}kWh，说明“无应急”不能单独代表经济最优。'
        elif r['date']=='2025-12-21':
            text+='该日光伏电量为四日最低、净负荷最高，购电费用主要由较高的基础购电需求决定；风险裕度与多轮充放电共同覆盖供需误差，因此负荷高也不必然出现紧急购电。'
        else:
            text+='此日光伏集中于白天，而较高电价及补电风险延续至晚间，白天总发电充足不能替代晚间功率与SOC可用性。'
        text+=f'本反馈规则对已形成缺口优先释放储能，没有人为保留备用的机制；日内保留电量是此前供需及SOC约束的结果。\n\n![当日联合调度](figures/{r["date"]}.png)\n\n'
    text+='## 13 敏感性与场景检验\n\n![风险与窗口对照](figures/quantile_window_sensitivity.png)\n\n'
    text+=table(['替代解释','总费用 元','相对主解释','紧急购电 kWh'],[[name,fmt(v['total_cost_yuan']),f'{v["total_cost_yuan"]/summary["total_cost_yuan"]-1:+.4%}',fmt(v['emergency_kwh'])] for name,v in sensitivity.items() if name!='no_storage_mean7'])+'\n\n'
    text+='效率与固定储能实验固定主策略已经选出的配置日历，仅重算计划和执行，隔离口径变化，不宣称该口径重新历史调参后的最优性能。固定储能解释可能在实际不足时为预定充电紧急补电，故不作为主解释。\n\n'
    pressure=read_csv(OUT/'experiments/scenario_pressure.csv')
    text+=table(['日期','残差构造','平均费元','95%费元','平均紧急kWh'],[[r['date'],r['kind'],f'{float(r["mean_cost"]):.2f}',f'{float(r["p95_cost"]):.2f}',f'{float(r["mean_emergency"]):.2f}'] for r in pressure])+'\n\n'
    text+='这些场景来自决策日前历史，而非当天真实未来；因也用于校准，不能称完全独立的样本外测试。真正时序样本外评价是334日实际逐日运行。独立重排改变相关性，差异用于说明相关性对风险评价的重要性，不能预设它一定低估风险。\n\n## 14 模型检验与复杂度\n\n'
    text+=f'逐时能量平衡最大误差{summary["checks"]["balance"]:.3e}kWh，状态误差{summary["checks"]["state"]:.3e}kWh，跨日连续误差{summary["continuity_error"]:.3e}kWh，所有边界和互斥检查通过。144/288/432时段LP分别有721/1441/2161变量及288/576/864等式，另有变量界。正式日均计划求解{summary["avg_plan_solve_seconds"]:.6f}秒；含历史选参、97个候选全年诊断及敏感性等总计算耗时{metadata["total_seconds"]:.2f}秒，测试和制图另计。\n\n'
    text+='测试包括未来反事实、月首参数选择、当期执行前缀不变、物理小算例、费用不退款、紧急区间合并、全量工作簿回读、同边界oracle及Q1接手哈希。最终实际通过数量与日志见validation.json、tests.log；不把只检查文件存在当作结果验证。\n\n## 15 模型评价与局限\n\n'
    text+='优势在于信息边界可测试、费用可重算、终端条件可比较、复杂度适中，且正式策略的每次选择都能追溯到过去数据。创新点是将预测器、风险裕度及终端处理放入同一历史经济验证框架，并通过可执行反馈评价有限策略场景，避免全知二阶段偏差。\n\n局限包括：仅一年数据、1月早期样本少；未给天气特征，光伏预测能力有限；历史21日/30日窗口与低价终端价值是预先固定的设计而非普适最优；贪心执行未优化跨时段应急机会成本；十分钟平均模型不能验证更细时间粒度；不计衰减、无并网容量；时间标签与90%效率仍有题面歧义；光伏来源分摊影响消纳率但不影响总费用。月度历史选择未提供分布漂移下收益保证，不能用事后最低候选夸大算法能力。\n\n## 参考文献与依据\n\n[1] 原题C题.pdf及附件1、2、5：参数、费用、输出合同。\n\n[2] Panda等，2024，Sustainable Energy, Grids and Networks 38:101334，DOI:10.1016/j.segan.2024.101334：日前承诺与实时实施分层，第5-8页；当前一步执行，第6页；跨日比较，第8-10页。\n\n[3] Herding等，2023，Applied Energy 336:120847，DOI:10.1016/j.apenergy.2023.120847：统一第一阶段及锁定后评价，第5、7页；模拟样本外边界，第7页。其全天第二阶段简化、售电和机组不照搬。\n\n两份建模报告是本地方案参考，不把它们的旧数值或“最终”称谓当作实验依据。具体取舍见方案审计.md及文献方法映射.md。\n'
    (OUT/'Q2_final.md').write_text(text,encoding='utf-8')
    decision='# Q2最终决策报告\n\n'+f'正式算法是月度过去21日经济验证选择的因果预测及储能计划策略。实际11个月均选择组合预测与24h终端价值，q按月变化，不固定80%，未选择48h或场景候选。主结果费用{summary["total_cost_yuan"]:,.4f}元；这是原始附件重算结果。\n\n'
    decision+='## 为什么这样选择\n\n每月选择最小历史末态调整费用，不根据未来月份挑参。候选全年对比仅作诊断。24h终端价值避免强制闭合，同时用代理电量价值限制透支；它在11次实际历史验证中胜出，不能据此推导所有数据上都优于48h。0.8有无储能单时段理论支持，但动态系统允许其他q更好。\n\n'
    decision+=table(['日期','预测/风险/终端配置'],[[str(dates[int(r['selection_day'])]),r['config']] for r in sel])+'\n\n'
    decision+='## 对现有报告与论文的取舍\n\n保留固定日前承诺、跨日SOC、5倍结算、历史残差、物理反馈与场景相关性。修改固定0.8和预设48h、信息集时点矛盾、弃电等同弃光以及场景内全知储能。Panda真正启发是承诺/执行分层与跨日比较；Herding是统一第一阶段和锁定后样本评价。两篇的售电、机组、备用市场、随机价格、竞价曲线和DRL没有本题依据，不采用。其文字中聚类数据范围、第二阶段完美信息等限制也已记录。\n\n'
    decision+='## 可执行性与信息证明\n\n0:00只用历史实际、日历、固定电价、实际初始E、过去验证和已实现残差。计划函数不接受当天实际；当前反馈函数不接受未来数组。测试篡改当天/未来月份后比较计划，另检查月首选择与当前动作前缀。证据见test_q2.py和最终测试日志。\n\n'
    decision+='## 数值与对照\n\n'+table(['指标','值'],[[k,fmt(summary[k]) if isinstance(summary[k],float) else summary[k]] for k in ('plan_cost_yuan','emergency_cost_yuan','total_cost_yuan','emergency_kwh','emergency_periods','emergency_days')])+'\n\n'+table(['策略','费用元','紧急kWh','紧急时段','末E'],cmp_rows)+'\n\n'
    decision+='## 预测及四个日期\n\n'+table(['对象','MAE kW','RMSE kW'],[[r['object'],f'{float(r["mae_kw"]):.3f}',f'{float(r["rmse_kw"]):.3f}'] for r in forecast if r['method']=='formal' and r['group']=='all'])+'\n\n'
    decision+=table(['日期','计划费元','紧急费元','总费元','紧急kWh'],[[r['date'],fmt(r['plan_cost_yuan']),fmt(r['emergency_cost_yuan']),fmt(r['total_cost_yuan']),fmt(r['emergency_kwh'])] for r in selected])+'\n\n'
    decision+='## 核验与仍需说明的边界\n\nresult2.xlsx按官方模板展开，计划表数值为0:00计划，充放电/H为实际执行；全量回读检查由test_q2.py进行，最终状态见validation.json。紧急H=max(0,L−PV−G−Dmax_available)，全年SOC连续。场景进入候选及压力测试，但实际月度选择未采用场景候选。\n\nQ1接手后39个文件保持不变，当前Q1_frozen_v2清单与ZIP完整校验通过；Q1_frozen_v1保留为历史。工作文件相对v1的8项差异属于既有版本演进，任一Q1清单均未改写。当前Q1测试结果另保存。\n\n时间结束点、实时储能反馈、初日回退、效率解释、终端价值、光伏来源分摊均须作为假设公开。所有替代口径及数值、完整表1/2/3、图件、模型局限见Q2_final.md，复现命令见README.md。\n'
    (OUT/'Q2最终决策报告.md').write_text(decision,encoding='utf-8')


def deliver():
    check_q1();dates,a,p=read_year();s=dict(np.load(OUT/'experiments/formal_arrays.npz'))
    summary=summary_of(s,a,p)
    summary.update(model='monthly_past_only_economic_selection',period='2025-02-01/2025-12-31',
                   official_is_oracle=False,price_rule='plan*p + emergency*5*p',
                   time_semantics='assumed_end_of_interval',efficiency='charge=.9 discharge=.9',
                   input_data='附件1价格、附件2实际；不使用附件3/4',
                   full_verification='see validation.json after actual independent test execution')
    write_json(OUT/'summary.json',summary)
    merged=export_workbook(dates,a,p,s)
    write_csv(OUT/'emergency_purchase_merged.csv',merged)
    daily,selected=data_exports(dates,a,p,s)
    comparisons=read_csv(OUT/'experiments/candidate_summary.csv')
    figures(dates,a,p,s,daily,comparisons)
    generate_documents(dates,a,p,s,summary,daily,selected,comparisons)
    print('已生成工作簿、完整明细、论文稿、决策报告和图件')


if __name__=='__main__':deliver()

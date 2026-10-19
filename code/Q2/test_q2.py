"""独立反事实/物理/结算/输出测试；不修改Q1和正式结果。"""
import sys,json,csv,hashlib,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import numpy as np
from openpyxl import load_workbook
from microgrid_common import Battery
from microgrid_online import read_year,plan_net,execute_day,feedback_step,merge_emergency,PERIODS,validate_execution,terminal_value
from forecast_q2 import build_features,forecast
from strategy_q2 import make_plan,Config,select_config,CANDIDATES
from solve_q2 import simulate,check_q1
Q2=Path(__file__).resolve().parent
HAS_Q2_OUTPUTS=(Q2/'experiments/formal_arrays.npz').exists() and (Q2/'summary.json').exists() and (Q2/'result2.xlsx').exists()


class Q2Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dates,cls.actual,cls.price=read_year()
        cls.features=build_features(cls.actual[:65])

    def test_dates_time_units(self):
        self.assertEqual(len(self.dates),365)
        self.assertEqual(len(self.dates[31:]),334)
        self.assertEqual(str(self.dates[31]),'2025-02-01')
        self.assertEqual(PERIODS[0],'0:00-0:10');self.assertEqual(PERIODS[-1],'23:50-24:00')
        self.assertEqual(PERIODS[60],'10:00-10:10')
        wb=load_workbook(ROOT/'附件/附件2.xlsx',read_only=True,data_only=True)
        self.assertAlmostEqual(wb.worksheets[0]['B2'].value/6,self.actual[0,0,0]);wb.close()

    def test_future_counterfactual_all_candidate_types(self):
        changed=self.actual[:65].copy();changed[40:,:,0]*=1.8;changed[40:,:,1]*=.1
        bank=build_features(changed)
        for conf in (Config('last',0,1),Config('blend',.7,1,'value'),Config('mean7',.9,3),Config(scenario=True)):
            a=make_plan(self.features[40],6000,self.price,conf)
            b=make_plan(bank[40],6000,self.price,conf)
            np.testing.assert_array_equal(a['G'],b['G'])
            x=execute_day(a['G'][:144],self.actual[40],6000)
            y=execute_day(b['G'][:144],changed[40],6000)
            self.assertGreater(abs(self.price@(x['H']-y['H'])),1)

    def test_future_month_does_not_change_earlier_features(self):
        changed=self.actual[:65].copy();changed[59:]*=20
        bank=build_features(changed)
        self.assertEqual(bank[31].history_sha256,self.features[31].history_sha256)
        for m in ('last','mean7','blend'):
            np.testing.assert_array_equal(bank[31].prediction[m],self.features[31].prediction[m])

    def test_residual_observation_date_and_horizon(self):
        for day,f in enumerate(self.features):
            self.assertLess(f.training_max,day)
            for method in ('last','mean7','blend'):
                for j in range(3):
                    ds=f.residual_dates[method][j]
                    self.assertTrue(all(k<day for k in ds))
                    for i,k in enumerate(ds):
                        origin=k-j
                        np.testing.assert_allclose(f.residual[method][j][i],self.actual[k]-self.features[origin].prediction[method][j])

    def test_month_selection_causality_and_whole_simulation(self):
        # 包含2月1日参数选择与后续执行，篡改当天和未来只能改变执行。
        candidates=(Config('last',0,1),Config('mean7',.8,2),Config('blend',.7,1,'value'))
        original,_=simulate(self.dates,self.actual,self.price,self.features,stop=33,candidates=candidates)
        changed=self.actual[:65].copy();changed[31:,:,0]*=1.5
        mutated,_=simulate(self.dates,changed,self.price,build_features(changed),stop=33,candidates=candidates)
        np.testing.assert_array_equal(original['G'][:32],mutated['G'][:32])
        self.assertEqual(original['config'][31],mutated['config'][31])
        self.assertGreater(np.abs(original['H'][31]-mutated['H'][31]).sum(),1)
        self.assertTrue(np.array_equal(original['E'][1:,0],original['E'][:-1,-1]))

    def test_execution_prefix_nonanticipativity(self):
        g=np.full(144,400.);a=self.actual[40].copy();b=a.copy();b[70:,0]*=3
        x=execute_day(g,a,6000);y=execute_day(g,b,6000)
        for k in ('C','D','H','W'):np.testing.assert_array_equal(x[k][:70],y[k][:70])
        np.testing.assert_array_equal(x['E'][:71],y['E'][:71])

    def test_feedback_bounds_and_no_mutation(self):
        g=np.array([0.,20000.,0.]);a=np.array([[10000.,0],[0,2000.],[10000.,0]])
        before=g.copy();s=execute_day(g,a,6000)
        np.testing.assert_array_equal(g,before)
        self.assertAlmostEqual(s['D'][0],5000/6)
        self.assertAlmostEqual(s['C'][1],5000/6)
        self.assertAlmostEqual(s['H'][0],10000-5000/6)
        validate_execution(a,s)

    def test_cost_five_times_and_no_refund(self):
        b=Battery();a=np.array([[100.,0.],[0.,0.]])
        s=execute_day(np.array([50.,100.]),a,1200,idle=True)
        p=np.array([2.,3.])
        self.assertEqual(p@s['G'],400.)
        self.assertEqual(5*p@s['H'],500.)
        self.assertEqual(s['unused_grid'][1],100.)

    def test_lp_milp_agree_and_mutex(self):
        records=[]
        for day,conf in [(2,Config()),(20,Config('blend',.75,1,'value')),(40,Config('mean7',.8,3))]:
            p=make_plan(self.features[day],6000,self.price,conf)
            m=plan_net(p['net'],np.tile(self.price,conf.horizon),6000,conf.terminal,integer=True)
            self.assertAlmostEqual(p['objective'],m['objective'],places=5)
            self.assertFalse(np.any((p['C']>1e-6)&(p['D']>1e-6)))
            records.append(p['objective']-m['objective'])

    def test_state_efficiency_hand_case(self):
        s=execute_day(np.array([100.,0.]),np.array([[0.,0.],[81.,0.]]),6000)
        np.testing.assert_allclose(s['C'],[100,0]);np.testing.assert_allclose(s['D'],[0,81])
        self.assertAlmostEqual(s['E'][-1],6000)

    def test_emergency_merge_no_midnight_join(self):
        h=np.zeros((2,144));h[0,60:63]=[1,2,3];h[0,-1]=4;h[1,0]=5
        r=merge_emergency(self.dates[:2],h)
        self.assertEqual([x['period'] for x in r],['10:00-10:30','23:50-24:00','0:00-0:10'])
        self.assertEqual([x['kwh'] for x in r],[6,4,5])

    def test_pv_attribution_bounds(self):
        s=execute_day(np.array([1000.]),np.array([[100.,50.]]),10800)
        self.assertEqual(s['W'][0],950);self.assertEqual(s['pv_curtail'][0],0)
        self.assertEqual(s['pv_curtail_upper'][0],50)

    def test_q1_protected(self):
        self.assertEqual(check_q1(),0)

    @unittest.skipUnless(HAS_Q2_OUTPUTS,'Run Q2 solve and report commands before output checks')
    def test_complete_saved_physics_and_cost(self):
        s=dict(np.load(Q2/'experiments/formal_arrays.npz'))
        for day in range(365):
            validate_execution(self.actual[day],{k:s[k][day] for k in ('G','C','D','H','W','E')})
        np.testing.assert_array_equal(s['E'][1:,0],s['E'][:-1,-1])
        self.assertGreater(np.unique(np.round(s['E'][31:,0],4)).size,1)
        summary=json.loads((Q2/'summary.json').read_text())
        self.assertAlmostEqual(float(((s['G'][31:]+5*s['H'][31:])@self.price).sum()),summary['total_cost_yuan'],places=5)
        self.assertFalse(any('oracle' in v for v in s['config'][31:]))

    @unittest.skipUnless(HAS_Q2_OUTPUTS,'Run Q2 solve and report commands before output checks')
    def test_excel_csv_full_readback(self):
        s=dict(np.load(Q2/'experiments/formal_arrays.npz'))
        wb=load_workbook(Q2/'result2.xlsx',read_only=True,data_only=True)
        self.assertEqual(wb.sheetnames,['计划购电量','充放电量','紧急购电量'])
        rows=list(wb.worksheets[0].values)
        self.assertEqual(len(rows),335);self.assertEqual(list(rows[0][1:145]),PERIODS)
        for i,row in enumerate(rows[1:],31):
            self.assertEqual(row[0].date(),self.dates[i])
            np.testing.assert_allclose(row[1:145],s['G'][i],atol=1e-8,rtol=0)
            self.assertAlmostEqual(row[145],s['G'][i].sum(),places=6)
            self.assertAlmostEqual(row[146],self.price@s['G'][i]+5*self.price@s['H'][i],places=6)
        battery=list(wb.worksheets[1].values)[1:]
        self.assertEqual(len(battery),334*6)
        for i in range(334):
            block=battery[i*6:(i+1)*6];d=i+31
            self.assertEqual(block[0][0].date(),self.dates[d])
            for j,row in enumerate(block):
                self.assertEqual(row[1],f'{j*4}:00-{(j+1)*4}:00')
                self.assertAlmostEqual(row[2],s['C'][d,j*24:(j+1)*24].sum(),places=6)
                self.assertAlmostEqual(row[3],s['D'][d,j*24:(j+1)*24].sum(),places=6)
            self.assertAlmostEqual(block[0][5],s['E'][d,0],places=6)
            self.assertAlmostEqual(block[1][5],s['E'][d,-1],places=6)
        erows=list(wb.worksheets[2].values)[1:];expected=merge_emergency(self.dates[31:],s['H'][31:])
        # 导出仅非零区间；无紧急日也有“无”占位。
        out=[];current=None
        for row in erows:
            if row[0] is not None:current=row[0].date().isoformat()
            if row[1]!='无':out.append((current,row[1],row[2]))
        self.assertEqual(len(out),len(expected))
        for x,y in zip(out,expected):
            self.assertEqual(x[:2],(y['date'],y['period']));self.assertAlmostEqual(x[2],y['kwh'],places=6)
        wb.close()
        with (Q2/'dispatch_10min.csv').open(encoding='utf-8-sig') as f:
            dispatch=list(csv.DictReader(f))
        self.assertEqual(len(dispatch),334*144)
        g=np.array([float(r['G_plan_kwh']) for r in dispatch]).reshape(334,144)
        h=np.array([float(r['H_real_kwh']) for r in dispatch]).reshape(334,144)
        np.testing.assert_allclose(g,s['G'][31:],atol=1e-8,rtol=0)
        np.testing.assert_allclose(h,s['H'][31:],atol=1e-8,rtol=0)

    @unittest.skipUnless(HAS_Q2_OUTPUTS,'Run Q2 solve and report commands before output checks')
    def test_selection_records_only_history(self):
        with (Q2/'experiments/monthly_selection.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
        dates=sorted(set(int(r['selection_day']) for r in rows))
        self.assertEqual(len(dates),11)
        for d in dates:
            group=[r for r in rows if int(r['selection_day'])==d]
            self.assertEqual(len(group),len(CANDIDATES))
            self.assertTrue(all(int(r['end'])==d-1 for r in group))
            chosen=[r for r in group if r['selected']=='True'];self.assertEqual(len(chosen),1)
            self.assertAlmostEqual(float(chosen[0]['score']),min(float(r['score']) for r in group),places=5)

    @unittest.skipUnless(HAS_Q2_OUTPUTS,'Run Q2 solve and report commands before output checks')
    def test_oracle_same_boundaries_lower_cost(self):
        oracle=dict(np.load(Q2/'experiments/oracle_arrays.npz'))
        formal=dict(np.load(Q2/'experiments/formal_arrays.npz'))
        self.assertAlmostEqual(oracle['E'][0],formal['E'][31,0],places=6)
        self.assertAlmostEqual(oracle['E'][-1],formal['E'][-1,-1],places=6)
        net=(self.actual[31:,:,0]-self.actual[31:,:,1]).ravel()
        np.testing.assert_allclose(oracle['G']+oracle['D']-oracle['C']-oracle['W'],net,atol=1e-5)
        self.assertLessEqual(np.tile(self.price,334)@oracle['G'],((formal['G'][31:]+5*formal['H'][31:])@self.price).sum()+1e-5)


if __name__=='__main__':unittest.main(verbosity=2)

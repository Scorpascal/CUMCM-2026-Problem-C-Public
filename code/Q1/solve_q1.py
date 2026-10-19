"""运行本文件可重建问题1的全部结果、图件、论文表格及检验记录。"""
from pathlib import Path
import csv
import hashlib
import json
import os
import sys
from dataclasses import asdict, replace

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
os.environ.setdefault('MPLCONFIGDIR', str(OUT / '.matplotlib'))
sys.path.insert(0, str(ROOT))
import numpy as np
import scipy
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from microgrid_common import Battery, clock_label, read_single_day, optimize, validate
from time_mapping import canonical_periods, period_bounds, validate_export_periods
from plot_preprocessing import plot_input_profiles


def csv_write(name, header, rows):
    with (OUT/name).open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f); writer.writerow(header); writer.writerows(rows)


def main():
    if any(OUT.glob('freeze_manifest*.json')):
        raise RuntimeError('Q1已冻结，禁止原地覆盖。请按冻结说明在独立解压副本中复现。')
    data = read_single_day(ROOT/'附件/附件1.xlsx')
    input_stats = plot_input_profiles(data, OUT)
    b = Battery()
    lp = optimize(data['price'],data['load'],data['pv'],b)
    s = optimize(data['price'],data['load'],data['pv'],b,integer=True)
    checks = validate(data,s,b)
    checks['milp_minus_lp_yuan'] = float(s['cost']-lp['cost'])
    if abs(checks['milp_minus_lp_yuan']) > 1e-5:
        raise AssertionError('MILP与LP下界不一致，需检查最优性')
    template = ROOT/'附件/附件5/result1.xlsx'
    template_wb = load_workbook(template, data_only=False)
    template_purchase = template_wb['计划购电量']
    template_labels = [template_purchase.cell(t,1).value for t in range(2,146)]
    template_headers = {
        ws.title: [ws.cell(1,c).value for c in range(1,ws.max_column+1)]
        for ws in template_wb
    }
    template_dimensions = {
        ws.title: (ws.max_row,ws.max_column)
        for ws in template_wb
    }
    template_styles = {
        ws.title: [[ws.cell(r,c).style_id for c in range(1,ws.max_column+1)]
                   for r in range(1,ws.max_row+1)]
        for ws in template_wb
    }
    template_wb.close()
    if len(template_labels) != 144 or len(set(template_labels)) != 144:
        raise AssertionError('官方模板的计划购电时间标签应为144个且不得重复')
    model_labels = canonical_periods(data['end_minutes'])
    template_offsets = [period_bounds(old)[0]-period_bounds(new)[0]
                        for old,new in zip(template_labels,model_labels)]
    if set(template_offsets) not in ({0}, {10}):
        raise ValueError('模板时间标签出现未知变化，需要重新核对，不能静默按行覆盖')
    for label,offset,end in zip(template_labels,template_offsets,data['end_minutes']):
        if period_bounds(label) != (int(end)-10+offset,int(end)+offset):
            raise ValueError('模板区间长度或结束时刻异常')
    base_grid = np.maximum(data['load']-data['pv'],0)
    base_cost = float(data['price']@base_grid)
    groups = [(f'{h}:00-{h+4}:00', float(s['charge'][h*6:(h+4)*6].sum()),
               float(s['discharge'][h*6:(h+4)*6].sum())) for h in range(0,24,4)]
    selected_indices = [60,72,84,96,108,120]
    selected = [(model_labels[t],float(s['grid'][t])) for t in selected_indices]
    rows = [[model_labels[t], data['price'][t],data['load'][t]/b.dt,data['pv'][t]/b.dt,
             s['grid'][t],s['charge'][t],s['discharge'][t],s['curtail'][t],
             s['energy'][t],s['energy'][t+1],data['price'][t]*s['grid'][t],base_grid[t]]
            for t in range(144)]
    header = ['时间段','电价_元每kWh','负载_kW','光伏预测_kW','购电_kWh','充电_母线侧_kWh',
              '放电_母线侧_kWh','弃光_kWh','期初储电_kWh','期末储电_kWh','购电费_元','无储能购电_kWh']
    csv_write('dispatch_10min.csv',header,rows)
    csv_write('table1.csv',['时间段','购电量_kWh'],selected+[('全天',float(s['grid'].sum()))])
    csv_write('table2.csv',['时间段','充电量_kWh','放电量_kWh'],groups)
    wb = load_workbook(template)
    ws = wb['计划购电量']
    for t in range(144):
        row = t+2
        if ws.cell(row,1).value != template_labels[t]:
            raise AssertionError('导出前模板时间标签发生变化')
        ws.cell(row,1,model_labels[t])
        ws.cell(row,2,float(s['grid'][t]))
    ws = wb['充放电量']
    for i,(_,c,d) in enumerate(groups,2):
        ws.cell(i,2,c); ws.cell(i,3,d)
    ws['E2']=float(s['energy'][0]); ws['E3']=float(s['energy'][-1])
    wb.save(OUT/'result1.xlsx')
    exported=load_workbook(OUT/'result1.xlsx',data_only=True)
    assert exported.sheetnames==['计划购电量','充放电量']
    assert {
        ws.title:(ws.max_row,ws.max_column) for ws in exported
    } == template_dimensions
    assert {
        ws.title:[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
        for ws in exported
    } == template_headers
    assert {
        ws.title:[[ws.cell(r,c).style_id for c in range(1,ws.max_column+1)]
                  for r in range(1,ws.max_row+1)]
        for ws in exported
    } == template_styles
    for t in range(144):
        row = t+2
        assert exported['计划购电量'].cell(row,1).value==model_labels[t]
        assert abs(exported['计划购电量'].cell(row,2).value-s['grid'][t])<1e-8
    validate_export_periods([exported['计划购电量'].cell(row,1).value
                            for row in range(2,146)],data['end_minutes'])
    for i,(_,c,d) in enumerate(groups,2):
        assert abs(exported['充放电量'].cell(i,2).value-c)<1e-8
        assert abs(exported['充放电量'].cell(i,3).value-d)<1e-8
    assert exported['充放电量']['E2'].value==exported['充放电量']['E3'].value==6000
    by_period = {exported['计划购电量'].cell(row,1).value:
                 exported['计划购电量'].cell(row,2).value for row in range(2,146)}
    for label,value in selected:
        assert abs(by_period[label]-value)<1e-8
    exported.close()
    csv_write('time_mapping.csv',['模型索引','附件1行号','附件1时刻_按结束点解释','输出行号',
              '原模板标签','模型自然日区间','输出Excel标签','原标签偏移_分钟','购电量_kWh'],
              [[t,t+2,clock_label(int(data['end_minutes'][t])),t+2,template_labels[t],
                model_labels[t],model_labels[t],template_offsets[t],float(s['grid'][t])] for t in range(144)])
    checks['template_match']='工作表、行列数、表头及样式沿用附件5；仅修正144个时间标签并填入所需数值'
    checks['time_label_match']='144个Excel标签逐项匹配模型及CSV，连续覆盖0:00至24:00；指定时段按标签查找核验'
    checks['unit_match']='购电量、充放电量和储电量均为kWh，购电费为元'
    checks['excel_export_order']='模型grid[0:144]按原始顺序逐项写入模板第2至145行，无循环平移'
    checks['table_indices']='表1索引为60、72、84、96、108、120；表2为六组连续24时段自然日汇总'
    checks['export_recheck']='144个购电量、6个指定时段、6组四小时充放电量及首末电量均已回读核对'
    alt_b=replace(b,eta_charge=np.sqrt(.9),eta_discharge=np.sqrt(.9))
    alt=optimize(data['price'],data['load'],data['pv'],alt_b,integer=True)
    validate(data,alt,alt_b)
    csv_write('efficiency_sensitivity.csv',['效率解释','充电效率','放电效率','购电费_元','购电量_kWh'],
              [('各90%',.9,.9,s['cost'],float(s['grid'].sum())),
               ('往返90%',alt_b.eta_charge,alt_b.eta_discharge,alt['cost'],float(alt['grid'].sum()))])
    metrics=dict(baseline_id='Q1_final',time_semantics='end_of_interval',
        model_period='0:00-24:00',excel_export_order=list(range(144)),cyclic_shift=False,
        official_template_labels_preserved=False,internal_checks_passed=True,
        unittest_execution='not_run_by_solve_q1; run Q1/test_q1.py separately',
        time_semantics_status='modeling_assumption_supported_by_data_and_report_not_explicit_in_problem',
        template_label_offset_minutes=sorted(set(template_offsets)),
        export_label_policy='canonical_natural_day_labels_no_value_shift',
        terminal_constraint_basis='问题1明确要求E144=E0；结合附录1的初始6000 kWh取E0=E144=6000',
        efficiency_interpretation='主模型将充放电效率90%解释为各90%；往返90%仅作敏感性比较',
        cost_yuan=s['cost'],grid_kwh=float(s['grid'].sum()),baseline_cost_yuan=base_cost,
        baseline_grid_kwh=float(base_grid.sum()),saving_yuan=base_cost-s['cost'],saving_rate=1-s['cost']/base_cost,
        load_kwh=float(data['load'].sum()),pv_kwh=float(data['pv'].sum()),
        curtail_kwh=float(s['curtail'].sum()),charge_kwh=float(s['charge'].sum()),discharge_kwh=float(s['discharge'].sum()),
        storage_loss_kwh=float(s['charge'].sum()-s['discharge'].sum()),
        energy_min_kwh=float(s['energy'].min()),energy_max_kwh=float(s['energy'].max()),
        initial_kwh=float(s['energy'][0]),terminal_kwh=float(s['energy'][-1]),
        lp_lower_bound_yuan=lp['cost'],mip_gap=s['mip_gap'],parameters=asdict(b),checks=checks,
        sensitivity_roundtrip_cost_yuan=alt['cost'],
        sources_sha256={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
                        for p in [ROOT/'附件/附件1.xlsx',template,ROOT/'C题.pdf',ROOT/'C题建模分析报告.docx']
                        if p.exists()},
        versions={'python':sys.version,'numpy':np.__version__,'scipy':scipy.__version__})
    (OUT/'summary.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding='utf-8')
    plot(data,s,b)
    report(metrics,selected,groups,header,rows,input_stats)
    audit_report(metrics)
    print(json.dumps({k:v for k,v in metrics.items() if k in ['cost_yuan','grid_kwh','saving_rate','checks']},ensure_ascii=False,indent=2))


def plot(data,s,b):
    x=np.arange(144)/6
    fig,axes=plt.subplots(4,1,figsize=(12,12),sharex=True,layout='constrained')
    axes[0].stairs(data['load']/b.dt,np.arange(145)/6,label='Load',color='#334155')
    axes[0].stairs(data['pv']/b.dt,np.arange(145)/6,label='PV forecast',color='#d99a13')
    axes[0].stairs(s['grid']/b.dt,np.arange(145)/6,label='Grid purchase',color='#2563eb')
    axes[0].set_ylabel('Power (kW)');axes[0].legend(ncol=3)
    axes[1].bar(x,s['charge']/b.dt,width=1/6,align='edge',label='Charge',color='#0d9488')
    axes[1].bar(x,-s['discharge']/b.dt,width=1/6,align='edge',label='Discharge',color='#ef7542')
    axes[1].set_ylabel('Battery power (kW)');axes[1].legend(ncol=2)
    axes[2].plot(np.arange(145)/6,s['energy'],color='#7c3aed')
    for e in [1200,6000,10800]:axes[2].axhline(e,color='#94a3b8',linestyle='--',linewidth=.8)
    axes[2].set_ylabel('Stored energy (kWh)');axes[2].set_ylim(0,12000)
    axes[3].stairs(data['price'],np.arange(145)/6,color='#be185d')
    axes[3].set_ylabel('Price (CNY/kWh)');axes[3].set_xlabel('Hour of day')
    for a in axes:a.grid(alpha=.2);a.set_xlim(0,24);a.set_xticks(np.arange(0,25,2))
    fig.suptitle('Q1: Optimal day-ahead microgrid dispatch',fontsize=16)
    fig.savefig(OUT/'dispatch.png',dpi=180);fig.savefig(OUT/'dispatch.svg');plt.close(fig)


def report(m,selected,groups,header,rows,input_stats):
    lines=['# 问题1：微电网确定性日前购电优化','',
        '> 基准版本：`Q1_final`。本文档、表格、图片、Excel结果和检验记录均由 `solve_q1.py` 从原始附件一键生成。','',
        '## 1 问题分析','',
        '问题1给定次日144个十分钟时段的电价、负荷与光伏预测，需要同时确定外网购电和储能充放电计划，使全天购电费用最低。决策的核心不是逐时段独立选择最低购电量，而是利用储能在时间维度转移电能：低价或光伏富余时充电，在高价时放电；同时满足功率平衡、储能效率、容量、功率、充放电互斥及日末电量回归约束。因此，该问题可建成确定性混合整数线性规划模型。','',
        '## 2 模型假设','',
        '1. 将一天均分为144个十分钟时段。原题未明示时间字段是起点还是终点；结合附件1从0:10至0:00+1的144条记录、自然日边界及建模报告，本文将时刻解释为区间结束点，并将功率视为区间平均功率，乘以1/6 h换算为电量。此为明确采用的离散化假设。',
        '2. 充电量与放电量均采用交流母线侧口径。附录1原文为“储能设备的充放电效率为90%”，未明示往返效率。本文采用充、放电效率各90%的主解释，往返效率为81%，另检验往返90%的解释。',
        '3. 微电网可从外网购电但不能向外网售电；题目未规定外网购电容量，故不人为添加购电上限。',
        '4. 允许弃光，不允许负荷削减；忽略储能自放电、寿命损耗费用及其他运维费用。',
        '5. 问题1明确要求0:00与24:00储电量相同，故日循环约束是题目硬约束。结合附录1给出的2025年1月1日0:00初始电量6000 kWh，本文以该初值求解，得到E₀=E₁₄₄=6000 kWh；并非仅为防止终端透支而自行添加末状态约束。','',
        '## 3 数据预处理','',
        '附件1包含144条十分钟记录。首先检查列名、记录数、时间连续性以及电价、负荷和光伏数据的有限性与非负性；随后将0:10至0:00+1解释为各区间结束时刻，构造0:00–0:10至23:50–24:00的自然日区间。负荷和光伏功率乘以1/6 h转换为优化所需的时段电量，电价保持元/kWh。','',
        '![图1 问题1三组输入数据的日内变化特征](input_daily_patterns.png)','',
        f"图1只使用附件1的三组输入数据。光伏在{input_stats['pv_max_period']}达到 {input_stats['pv_max_kw']:.1f} kW；负荷在白天保持较高水平，最大值为 {input_stats['load_max_kw']:.1f} kW；最低电价 {input_stats['price_min_yuan_per_kwh']:.4f} 元/kWh 出现在{input_stats['price_min_period']}。三者峰谷在时间上明显错开。",'',
        '## 4 储能“谷充峰放”的数据依据','',
        '夜间光伏为零但电价较低，白天负荷处于高平台且光伏在正午达到峰值，傍晚至晚间电价明显升高。由此，储能可以在夜间低价时段或正午光伏富余时吸收电能，并在晚间高价时段释放电能。该观察仅用于解释优化机制，具体充放电量仍由后续约束模型统一决定。','',
        '## 5 符号说明','',
        '|符号|含义|单位|',
        '|---|---|---|',
        '|$t$|十分钟时段编号，$t=0,1,\\ldots,143$|—|',
        '|$\\lambda_t$|时段电价|元/kWh|',
        '|$L_t,P_t$|负荷电量、光伏预测电量|kWh|',
        '|$G_t,C_t,D_t,W_t$|购电量、母线侧充电量、母线侧放电量、弃光量|kWh|',
        '|$E_t$|时段$t$开始时的储电量|kWh|',
        '|$z_t$|充电状态变量，充电为1、放电为0|—|',
        '|$\\eta_c,\\eta_d$|充电效率、放电效率，均取0.9|—|',
        '|$\\Delta t$|时段长度，取$1/6$|h|','',
        '## 6 模型建立','',
        '### 6.1 目标函数','',
        '以全天从外网购电的实际费用最小为目标，不额外添加会改变原问题最优解的储能吞吐惩罚：','',
        '$$\\min F=\\sum_{t=0}^{143}\\lambda_tG_t.\\tag{1}$$','',
        '### 6.2 功率平衡约束','',
        '每个时段内，外网购电、光伏和储能放电共同满足负荷、储能充电与弃光需求：','',
        '$$G_t+P_t+D_t=L_t+C_t+W_t,\\quad \\forall t.\\tag{2}$$','',
        '### 6.3 储能状态方程','',
        '考虑充放电效率后，储电量按下式连续演化：','',
        '$$E_{t+1}=E_t+\\eta_cC_t-\\frac{D_t}{\\eta_d},\\quad \\forall t.\\tag{3}$$','',
        '### 6.4 SOC、功率和互斥约束','',
        '储能容量为12000 kWh，SOC允许范围为10%至90%。问题1规定首末储电量相同，结合附录1初始值6000 kWh，得到','',
        '$$1200\\le E_t\\le10800,\\quad E_0=E_{144}=6000.\\tag{4}$$','',
        '最大充放电功率为5000 kW。用二进制变量显式禁止同一十分钟时段同时充、放电：','',
        '$$0\\le C_t\\le5000\\Delta t\\,z_t,\\quad '
        '0\\le D_t\\le5000\\Delta t(1-z_t),\\quad z_t\\in\\{0,1\\}.\\tag{5}$$','',
        '购电和弃光边界为','',
        '$$G_t\\ge0,\\quad0\\le W_t\\le P_t.\\tag{6}$$','',
        '## 7 求解方法','',
        '采用 SciPy 调用 HiGHS 求解上述MILP；另去掉式(5)中的二进制互斥变量，求解连续LP作为成本下界。MILP目标值与LP下界在数值容差内一致，且导出结果通过逐时段功率平衡、状态递推、容量、功率、互斥和Excel回读检查，因此本数据下获得全局最优方案。',
        f"LP下界为 {m['lp_lower_bound_yuan']:.8f} 元，MILP最优费用为 {m['cost_yuan']:.8f} 元，相对MIP间隙为 {m['mip_gap']:.3g}。",'',
        '## 8 求解结果','',
        '### 表1 指定时段购电量与全天汇总','',
        '|时间段|购电量（kWh）|时间段|购电量（kWh）|时间段|购电量（kWh）|',
        '|---|---:|---|---:|---|---:|']
    for k in (0,3): lines.append('|'+'|'.join(f'{a}|{v:.6f}' for a,v in selected[k:k+3])+'|')
    lines += [f"|全天购电量|{m['grid_kwh']:.6f}|全天购电费（元）|{m['cost_yuan']:.6f}|||",'',
        '### 表2 储能充放电量与首末储电量','',
        '|时间段|充电量（kWh）|放电量（kWh）|时间段|充电量（kWh）|放电量（kWh）|',
        '|---|---:|---:|---|---:|---:|']
    for k in (0,2,4):lines.append('|'+'|'.join(f'{a}|{c:.6f}|{d:.6f}' for a,c,d in groups[k:k+2])+'|')
    lines += ['|0:00储电量|6000||24:00储电量|6000||','',
        '![图2 负荷、光伏、购电、储能与电价联合调度结果](dispatch.png)','',
        '图2给出了同一时间轴上的四联图。调度策略与电价、光伏及SOC边界一致：凌晨电价较低，系统在满足负荷的同时以最高5000 kW功率阶段性充电；清晨电价上升后储能放电，6:00—9:00附近显著减少外网购电；中午光伏出力达到高位，系统优先消纳光伏并将可用余量存入储能；18:00—21:00处于全天高价区间，储能集中放电，其中19时附近外网购电降为0；21时以后电价下降，系统重新充电，使日末储电量严格回到6000 kWh。','',
        '### 8.1 经济性分析','',
        f"无储能基准按 max(负载−光伏,0) 购电，费用为 {m['baseline_cost_yuan']:.6f} 元；优化节省 {m['saving_yuan']:.6f} 元，节省率 {m['saving_rate']:.4%}。",
        f"全天负载 {m['load_kwh']:.6f} kWh，光伏 {m['pv_kwh']:.6f} kWh，弃光 {m['curtail_kwh']:.6f} kWh。",
        f"母线侧总充电 {m['charge_kwh']:.6f} kWh，总放电 {m['discharge_kwh']:.6f} kWh，储能损耗 {m['storage_loss_kwh']:.6f} kWh。储电量范围 {m['energy_min_kwh']:.6f}–{m['energy_max_kwh']:.6f} kWh。",
        '全天电量可通过“购电+光伏=负载+弃光+储能损耗”复核。费用下降主要来自储能跨时段价差套利与光伏就地消纳，而非减少必须供给的负荷。','',
        '### 8.2 三项边界说明','',
        '1. **高购电尖峰没有越界。** 购电量由式(2)决定。在低价充电时，外网需同时承担小区净负荷和储能充电功率，因此图中的购电尖峰约等于“负荷减光伏+5000 kW最大充电功率”。5000 kW限制约束的是储能充电功率，不是外网总购电功率。',
        '2. **四小时汇总表中充、放电量均为正不代表同时充放电。** 表2汇总了一个四小时窗口内24个十分钟决策；不同子时段可先充后放。式(5)在每个十分钟时段强制互斥，校验结果中的同时充放电时段数为0。',
        '3. **外网购电不设人为上限。** 题目仅给出储能容量和充放电功率限制，没有给出并网点容量。若擅自增加购电上限会改变可行域和最优费用，因此主模型只保留非负购电约束；如后续题目补充并网容量，再将其作为场景参数加入。','',
        '## 9 模型检验与敏感性分析','',
        '### 9.1 效率口径敏感性','',
        f"若将题目90%解释为往返效率，充放电效率各取√0.9，最优费用为 {m['sensitivity_roundtrip_cost_yuan']:.6f} 元。该结果仅作敏感性比较，result1.xlsx 采用各90%的主模型。",'',
        '### 9.2 官方模板与结果一致性校验','',
        '原题附录2仅明确0:00+1等于当天24:00，没有明文规定所有记录均是区间结束点。本文根据数据序列和自然日要求采用结束点解释。原模板标签从0:10–0:20至0:00+1–0:10+1，按字面缺少当天首十分钟且多出次日首十分钟；它与本文采用的区间定义相差10分钟。',
        '本次输出保留两张工作表、行列数、表头及样式，仅把计划购电量的时间列修正为0:00–0:10至23:50–24:00，购电数值不循环平移。第一行0:00–0:10对应1406.6411 kWh，Excel、CSV、论文按同一自然日区间解释。time_mapping.csv保留原模板标签、输入行号、模型区间、输出标签和购电量的完整逐行对照。原始附件不修改。',
        '该选择消除了交付文件内部的标签与数据矛盾，但现有本地资料不能确认外部评分程序的时间约定，也不构成官方勘误。若获得官方明确说明，应依据说明统一调整输入映射、统计和输出，不能只循环移动购电列。',
        '购电量、充放电量和储电量均以kWh填写，购电费以元计。程序保存后重新打开工作簿，逐项核对144个购电量、6个指定时段、6组四小时充放电累计值以及0:00和24:00储电量。',
        'summary.json记录输入文件SHA256、运行版本、约束误差、LP下界、模板一致性和Excel回读校验。dispatch_10min.csv保留未人为舍入的完整决策；表格显示小数不影响内部精度。',
        'summary.json中的internal_checks_passed仅表示生成脚本内置检查通过。生成脚本不自动执行unittest；回归测试需另行运行test_q1.py，不再写入未经本次测试证明的tests_passed字段。',
        '表1/表2以本报告为论文排版版本，CSV用于进一步排版。', '',
        '## 10 模型评价','',
        '模型结构简洁、物理约束完整，并通过MILP与LP下界对照证明本算例最优性；同时保留十分钟粒度结果，避免四小时汇总掩盖时序行为。其局限在于采用确定性预测并忽略储能衰减和并网容量。问题2—4若引入预测误差、实时修正或新设备，应在本公共储能物理模型之上扩展信息与费用逻辑，而不改变问题1的 `Q1_final` 基准。','',
        '## 参考资料','',
        '- C题.pdf：问题1、附录1、附录2及表1/表2。',
        '- C题建模分析报告.docx：第2至4节及附录假设。',
        '- Silva等（2020），Optimal Day-Ahead Scheduling of Microgrids with Battery Energy Storage System，第3.1.3–3.1.4节：平衡、效率与储能约束。',
        '- Šimić等（2022），Optimal Scheduling of Battery Energy Storage in Microgrid to Minimize Electricity and Fuel Costs，第III节：互斥和日初日末相等。',
        '- 参考资料/问题1相关论文/文献目录与摘录.md：采用购电成本目标，不照搬售电、柴油机或寿命成本。']
    text = '\n'.join(lines)+'\n'
    (OUT/'Q1_final.md').write_text(text,encoding='utf-8')
    (OUT/'问题1结果与模型说明.md').write_text(text,encoding='utf-8')


def audit_report(m):
    text = f'''# 问题1口径核验

本说明由 solve_q1.py 生成，依据本地 C题.pdf、附件1、附件5及建模分析报告。事实、建模解释和无法确认的评分约定分开记录。

## 1 十分钟时间映射

- 原题第2页附录2明文：0:00+1表示次日0:00，即当天24:00。未明文说明功率记录是区间起点、终点还是瞬时采样。
- 附件1 A2:A145共144条记录，为0:10至0:00+1；建模报告第2.1节建议按结束点解释。
- 原题第1页要求当天计划和0:00/24:00边界，表2覆盖0:00至24:00。结束点及区间平均功率解释与之相容，本文继续采用该假设。
- 原模板A2为0:10-0:20，A145为0:00+1-0:10+1。按字面整体覆盖0:10至次日0:10，无法原样同时表达本文自然日首末时段。
- 旧输出虽然按行号对应，但标签和数值不一致。新输出修正时间标签为0:00-0:10至23:50-24:00，保留数据顺序和模板结构、样式。详细证据链见time_mapping.csv。
- 不循环平移：移动购电量而不同步负载、光伏、价格、储能状态会破坏物理对应；当前也没有足够依据把当天策略延伸到次日首十分钟。
- 尚未取得官方勘误或评分程序说明。修正标签是本地交付的一致性选择，不声称已消除外部评分约定的不确定性。

## 2 效率

原题第2页附录1原文：“储能设备的充放电效率为90%”。它没有使用“综合效率”“往返效率”，也没有分别列出两个效率。主模型按各90%处理，与建模报告第3.2节建议一致；此为采用的解释，不能表述成原题分别明示两个90%。

|解释|充电效率|放电效率|往返效率|购电费（元）|
|---|---:|---:|---:|---:|
|主模型，各90%|0.9|0.9|0.81|{m['cost_yuan']:.6f}|
|敏感性，往返90%|√0.9|√0.9|0.9|{m['sensitivity_roundtrip_cost_yuan']:.6f}|

两种解释成本相差 {m['cost_yuan']-m['sensitivity_roundtrip_cost_yuan']:.6f} 元。result1.xlsx只使用主模型，敏感性结果不混入提交列。

## 3 日末电量

原题第1页问题1明文：“储能设备在0:00和24:00的储电量相同”。因此E144=E0是硬约束。

附录1给出2025年1月1日0:00电量6000 kWh。以该初始状态和题目指定的循环边界结合，取E0=E144=6000 kWh。原题不是在同一句中直接写出两个6000；论文已说明两项依据。仅因避免终端透支而将此列为额外假设的表述不准确。

## 4 检查状态及复现

summary.json改用internal_checks_passed，含义仅限物理约束、最优性和生成文件内置回读检查通过；unittest_execution明确注明生成器未运行独立测试。

```bash
.venv/bin/python Q1/solve_q1.py
.venv/bin/python Q1/test_q1.py
```

回归测试按区间标签检查Excel/CSV/表1数值对应，并包含故意偏移或循环移动标签的反例。仅“工作表结构与原模板一致”不再被视为时间正确的证明。
'''
    (OUT/'口径核验.md').write_text(text,encoding='utf-8')


if __name__=='__main__':main()

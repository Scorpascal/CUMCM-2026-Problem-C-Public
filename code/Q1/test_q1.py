"""可手算的小算例，检验效率、互斥、功率限值、弃光和时间解析。"""
import sys
import unittest
from pathlib import Path
from dataclasses import replace
from datetime import time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import numpy as np
from microgrid_common import Battery, optimize, validate, end_minutes
from time_mapping import canonical_periods, validate_export_periods
Q1=Path(__file__).resolve().parent
HAS_Q1_OUTPUTS=(Q1/'result1.xlsx').exists() and (Q1/'summary.json').exists()


class DispatchTests(unittest.TestCase):
    def solve(self,price,load,pv,b):
        data={k:np.array(v,float) for k,v in zip(['price','load','pv'],[price,load,pv])}
        s=optimize(**data,battery=b,initial=0,terminal=0,integer=True)
        validate(data,s,b,initial=0,terminal=0)
        return s

    def test_efficiency_and_arbitrage(self):
        b=Battery(capacity=10,minimum=0,maximum=10,power=20,dt=1,eta_charge=.9,eta_discharge=.9)
        s=self.solve([1,3],[0,8.1],[0,0],b)
        self.assertAlmostEqual(s['cost'],10)
        np.testing.assert_allclose(s['charge'],[10,0],atol=1e-7)
        np.testing.assert_allclose(s['discharge'],[0,8.1],atol=1e-7)

    def test_power_limit(self):
        b=Battery(capacity=10,minimum=0,maximum=10,power=5,dt=1,eta_charge=1,eta_discharge=1)
        s=self.solve([1,3],[0,8],[0,0],b)
        self.assertAlmostEqual(s['cost'],14)

    def test_curtailment_and_no_sale(self):
        b=Battery(capacity=10,minimum=0,maximum=10,power=5,dt=1)
        s=self.solve([1],[0],[20],b)
        self.assertAlmostEqual(s['cost'],0)
        self.assertAlmostEqual(s['curtail'][0],20)

    def test_equal_prices_do_not_cycle(self):
        b=Battery(capacity=10,minimum=0,maximum=10,power=5,dt=1)
        s=self.solve([1,1],[4,4],[0,0],b)
        self.assertAlmostEqual(s['cost'],8)
        self.assertAlmostEqual(s['charge'].sum(),0)

    def test_time_formats(self):
        self.assertEqual(end_minutes(time(0,10)),10)
        self.assertEqual(end_minutes('10:10'),610)
        self.assertEqual(end_minutes('0:00+1'),1440)

    @unittest.skipUnless(HAS_Q1_OUTPUTS,'Run Q1/solve_q1.py before output checks')
    def test_generated_final_report_contract(self):
        report = Path(__file__).with_name('Q1_final.md')
        self.assertTrue(report.exists())
        text = report.read_text(encoding='utf-8')
        for required in (
            '基准版本：`Q1_final`',
            '## 1 问题分析',
            '## 2 模型假设',
            '## 3 数据预处理',
            'input_daily_patterns.png',
            '## 4 储能“谷充峰放”的数据依据',
            '## 5 符号说明',
            '\\tag{1}',
            '\\tag{6}',
            '高购电尖峰没有越界',
            '四小时汇总表中充、放电量均为正不代表同时充放电',
            '外网购电不设人为上限',
        ):
            self.assertIn(required, text)
        self.assertEqual(
            text,
            Path(__file__).with_name('问题1结果与模型说明.md').read_text(encoding='utf-8'),
        )

    @unittest.skipUnless(HAS_Q1_OUTPUTS,'Run Q1/solve_q1.py before output checks')
    def test_preprocessing_figure_contains_inputs_only(self):
        import json
        q1=Path(__file__).resolve().parent
        for name in ('input_daily_patterns.png','input_daily_patterns.svg','input_statistics.json'):
            self.assertTrue((q1/name).is_file())
            self.assertGreater((q1/name).stat().st_size,100)
        svg=(q1/'input_daily_patterns.svg').read_text(encoding='utf-8')
        for required in ('小区负荷','光伏预测','外网电价'):
            self.assertIn(required,svg)
        for forbidden in ('购电','充电','放电','储电','Grid purchase','Battery'):
            self.assertNotIn(forbidden,svg)
        stats=json.loads((q1/'input_statistics.json').read_text(encoding='utf-8'))
        self.assertEqual(stats['record_count'],144)
        self.assertEqual(stats['period_start'],'0:00-0:10')
        self.assertEqual(stats['period_end'],'23:50-24:00')

    @unittest.skipUnless(HAS_Q1_OUTPUTS,'Run Q1/solve_q1.py before output checks')
    def test_official_workbook_contract(self):
        from openpyxl import load_workbook
        q1 = Path(__file__).resolve().parent
        template = q1.parent/'附件/附件5/result1.xlsx'
        source = load_workbook(template,data_only=False)
        result = load_workbook(q1/'result1.xlsx',data_only=True)
        self.assertEqual(source.sheetnames,result.sheetnames)
        for source_sheet,result_sheet in zip(source,result):
            self.assertEqual(
                (source_sheet.max_row,source_sheet.max_column),
                (result_sheet.max_row,result_sheet.max_column),
            )
            self.assertEqual(
                [source_sheet.cell(1,c).value for c in range(1,source_sheet.max_column+1)],
                [result_sheet.cell(1,c).value for c in range(1,result_sheet.max_column+1)],
            )
            self.assertEqual(
                [[source_sheet.cell(r,c).style_id
                  for c in range(1,source_sheet.max_column+1)]
                 for r in range(1,source_sheet.max_row+1)],
                [[result_sheet.cell(r,c).style_id
                  for c in range(1,result_sheet.max_column+1)]
                 for r in range(1,result_sheet.max_row+1)],
            )
        source_labels = [source['计划购电量'].cell(r,1).value for r in range(2,146)]
        result_labels = [result['计划购电量'].cell(r,1).value for r in range(2,146)]
        validate_export_periods(result_labels,range(10,1441,10))
        self.assertEqual(result_labels[0],'0:00-0:10')
        self.assertEqual(result_labels[-1],'23:50-24:00')
        self.assertEqual(source_labels[0],'0:10-0:20')
        self.assertEqual(source_labels[-1],'0:00+1-0:10+1')
        self.assertTrue(all(
            isinstance(result['计划购电量'].cell(r,2).value,(int,float))
            for r in range(2,146)
        ))
        source.close()
        result.close()

    @unittest.skipUnless(HAS_Q1_OUTPUTS,'Run Q1/solve_q1.py before output checks')
    def test_final_natural_day_mapping_and_metrics(self):
        import csv
        import json
        from openpyxl import load_workbook

        q1 = Path(__file__).resolve().parent
        data = __import__('microgrid_common').read_single_day(q1.parent/'附件/附件1.xlsx')
        battery = Battery()
        milp = optimize(data['price'],data['load'],data['pv'],battery,integer=True)
        lp = optimize(data['price'],data['load'],data['pv'],battery)
        checks = validate(data,milp,battery)
        self.assertEqual(len(milp['grid']),144)
        self.assertAlmostEqual(milp['cost'],lp['cost'],places=5)
        self.assertEqual(milp['mip_gap'],0)
        self.assertAlmostEqual(milp['cost'],35126.948589289634,places=5)
        self.assertAlmostEqual(milp['grid'].sum(),59482.69899835391,places=5)
        self.assertAlmostEqual(milp['energy'][0],6000)
        self.assertAlmostEqual(milp['energy'][-1],6000)
        self.assertGreaterEqual(milp['energy'].min(),1200-1e-6)
        self.assertLessEqual(milp['energy'].max(),10800+1e-6)
        self.assertEqual(checks['simultaneous_charge_discharge_count'],0)
        self.assertLessEqual(checks['power_bound_violation_kw'],1e-6)
        self.assertLessEqual(checks['balance_max_error_kwh'],1e-6)
        self.assertLessEqual(checks['state_max_error_kwh'],1e-6)

        workbook = load_workbook(q1/'result1.xlsx',read_only=True,data_only=True)
        exported = np.array([
            workbook['计划购电量'].cell(row,2).value for row in range(2,146)
        ])
        np.testing.assert_allclose(exported,milp['grid'],atol=1e-8)
        self.assertAlmostEqual(exported[0],milp['grid'][0])
        self.assertAlmostEqual(exported[-1],milp['grid'][143])
        workbook.close()

        with (q1/'table1.csv').open(encoding='utf-8-sig') as f:
            table1 = list(csv.DictReader(f))
        indices = [60,72,84,96,108,120]
        self.assertEqual(
            [float(row['购电量_kWh']) for row in table1[:-1]],
            [float(milp['grid'][t]) for t in indices],
        )
        with (q1/'table2.csv').open(encoding='utf-8-sig') as f:
            table2 = list(csv.DictReader(f))
        for group,row in enumerate(table2):
            start = group*24
            stop = start+24
            self.assertAlmostEqual(float(row['充电量_kWh']),milp['charge'][start:stop].sum())
            self.assertAlmostEqual(float(row['放电量_kWh']),milp['discharge'][start:stop].sum())

        summary = json.loads((q1/'summary.json').read_text(encoding='utf-8'))
        self.assertEqual(summary['time_semantics'],'end_of_interval')
        self.assertEqual(summary['model_period'],'0:00-24:00')
        self.assertEqual(summary['excel_export_order'],list(range(144)))
        self.assertFalse(summary['cyclic_shift'])
        self.assertFalse(summary['official_template_labels_preserved'])
        self.assertTrue(summary['internal_checks_passed'])
        self.assertNotIn('tests_passed',summary)
        self.assertEqual(summary['unittest_execution'],
                         'not_run_by_solve_q1; run Q1/test_q1.py separately')

    def test_shifted_labels_are_rejected(self):
        # 旧版首行标签错误即使数值未平移，也必须判定为不一致。
        from microgrid_common import clock_label
        labels=[f'{clock_label(t)}-{clock_label(t+10)}' for t in range(10,1450,10)]
        with self.assertRaises(ValueError):
            validate_export_periods(labels,range(10,1441,10))

    def test_rotated_labels_are_rejected(self):
        labels=canonical_periods(range(10,1441,10))
        with self.assertRaises(ValueError):
            validate_export_periods(labels[1:]+labels[:1],range(10,1441,10))

    @unittest.skipUnless(HAS_Q1_OUTPUTS,'Run Q1/solve_q1.py before output checks')
    def test_excel_csv_table1_agree_by_time_label(self):
        import csv
        from openpyxl import load_workbook
        q1=Path(__file__).resolve().parent
        with (q1/'dispatch_10min.csv').open(encoding='utf-8-sig') as f:
            dispatch=list(csv.DictReader(f))
        with (q1/'table1.csv').open(encoding='utf-8-sig') as f:
            table1=list(csv.DictReader(f))
        with (q1/'time_mapping.csv').open(encoding='utf-8-sig') as f:
            mapping=list(csv.DictReader(f))
        w=load_workbook(q1/'result1.xlsx',data_only=True)
        actual={label:value for label,value in list(w['计划购电量'].values)[1:]}
        expected={row['时间段']:float(row['购电_kWh']) for row in dispatch}
        self.assertEqual(set(actual),set(expected))
        for label in expected:
            self.assertAlmostEqual(actual[label],expected[label],places=8)
        for row in table1[:-1]:
            self.assertAlmostEqual(actual[row['时间段']],float(row['购电量_kWh']),places=8)
        self.assertEqual(len(mapping),144)
        for index,row in enumerate(mapping):
            self.assertEqual(int(row['模型索引']),index)
            self.assertEqual(int(row['附件1行号']),index+2)
            self.assertEqual(row['模型自然日区间'],row['输出Excel标签'])
            self.assertAlmostEqual(actual[row['输出Excel标签']],float(row['购电量_kWh']),places=8)
        w.close()


if __name__=='__main__':unittest.main()

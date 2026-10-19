"""独立重算、手算、未来预报/实际反事实和全量交付回读。"""
import sys,json,csv,unittest
from pathlib import Path
from datetime import datetime,timedelta
from dataclasses import replace
ROOT=Path(__file__).resolve().parents[1];OUT=Path(__file__).resolve().parent;sys.path.insert(0,str(ROOT))
import numpy as np
from openpyxl import load_workbook
from microgrid_forecasts import read_pv_forecasts,interpolate_pv,released_curve
from microgrid_adjustment import settlement,adjustment_plan
from microgrid_online import read_year,execute_day,validate_execution,PERIODS,merge_emergency,terminal_value
from q3_forecast import FeatureFactory,StageFeatures
from q3_model import Config,choose_stage,run_day,simulate_paths,CANDIDATES
from solve_q3 import check_protected

class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dates,cls.a,cls.p=read_year();cls.r=read_pv_forecasts(ROOT/'附件/附件3.xlsx');cls.f=FeatureFactory(cls.a,cls.r)

    def test_forecast_keys_dates_and_units(self):
        self.assertEqual(len(self.r),1460)
        for issue,rec in self.r.items():
            self.assertEqual(rec.targets[0],issue+timedelta(hours=1));self.assertEqual(rec.targets[-1],issue+timedelta(hours=24))
        self.assertEqual(self.r[datetime(2025,1,1,18)].targets[-1],datetime(2025,1,2,18))
        self.assertEqual(PERIODS[36],'6:00-6:10');self.assertEqual(PERIODS[108],'18:00-18:10')
        rec=self.r[datetime(2025,3,20,6)]
        for method in ('constant','linear','pchip'):
            v=interpolate_pv(rec,10.,method)
            self.assertEqual(len(v),144);np.testing.assert_allclose(v[5::6]/6,np.array(rec.power)/6,atol=1e-8)
            self.assertGreaterEqual(v.min(),0)

    def test_reject_future_release(self):
        now=datetime(2025,3,20,6)
        with self.assertRaises(AssertionError):released_curve([self.r[now+timedelta(hours=6)]],now,0)

    def test_residual_targets_have_all_arrived(self):
        for day in (1,2,10,31,100):
            for k in range(4):
                f=self.f.feature(day,k)
                self.assertLessEqual(f.residual_latest_end,f.now)
                for j,d in enumerate(f.residual_origins):
                    lo=d*144+k*36
                    np.testing.assert_allclose(f.residual[j],self.f.flat[lo:lo+288]-self.f.raw(d,k,'linear',0.,k))

    def test_publication_time_counterfactual(self):
        day=60;base=datetime(2025,1,1)+timedelta(days=day)
        for cutoff in (1,2,3):
            records=dict(self.r)
            for k in range(cutoff,4):
                dt=base+timedelta(hours=6*k);rec=records[dt];records[dt]=replace(rec,power=tuple(np.array(rec.power)*.01+5000))
            x,_=run_day(self.f,day,6000,self.p,Config())
            y,_=run_day(FeatureFactory(self.a,records),day,6000,self.p,Config())
            np.testing.assert_array_equal(x['versions'][:cutoff],y['versions'][:cutoff])
            np.testing.assert_array_equal(x['G'][:cutoff*36],y['G'][:cutoff*36])

    def test_actual_future_counterfactual_all_stage_decisions(self):
        day=60
        for cutoff in (0,36,72,108):
            a=self.a.copy();a[day,cutoff:,0]*=2;a[day,cutoff:,1]*=.01;a[day+1:]*=1.7
            x,_=run_day(self.f,day,6000,self.p,Config(gain=.5,bias=True))
            y,_=run_day(FeatureFactory(a,self.r),day,6000,self.p,Config(gain=.5,bias=True))
            np.testing.assert_array_equal(x['versions'][:cutoff//36+1],y['versions'][:cutoff//36+1])
            for key in ('C','D','H','W'):np.testing.assert_array_equal(x[key][:cutoff],y[key][:cutoff])
            self.assertGreater(np.abs(x['H'][cutoff:]-y['H'][cutoff:]).sum(),1.)

    def test_nonretroactivity_and_adjustment_identity(self):
        for rule in ('A','B','B_net'):
            s,_=run_day(self.f,80,6000,self.p,Config(),rule=rule)
            for k in (1,2,3):
                np.testing.assert_array_equal(s['versions'][k,:k*36],s['versions'][k-1,:k*36])
                np.testing.assert_allclose(s['versions'][k],s['versions'][k-1]+s['U'][k-1]-s['R'][k-1],atol=1e-8)
            self.assertFalse(np.any((s['U']>1e-7)&(s['R']>1e-7)))

    def test_settlement_hand_multiple_transactions(self):
        for rule,expected in [('A',165),('B',125),('B_net',95)]:
            g=np.array([100.]);base=g.copy();total=100.
            for new in (80.,110.,90.):
                bill=settlement(g,np.array([new]),np.ones(1),rule,base);total+=bill['charge'].sum();g[:]=new
            self.assertAlmostEqual(total,expected)
        for rule,expected in [('A',140),('B',120),('B_net',100)]:
            total=100.;g=np.array([100.]);base=g.copy()
            for new in (120.,100.):
                total+=settlement(g,np.array([new]),np.ones(1),rule,base)['charge'].sum();g[:]=new
            self.assertAlmostEqual(total,expected)

    def synthetic(self,pv=0.,load=100.,initial=1200.,old=50.,rule='A',threshold=1e-6):
        pred=np.column_stack([np.full(288,load),np.full(288,pv)])
        f=StageFeatures(36,36,pred,np.zeros((0,288,2)),(),-1,np.zeros(288))
        return choose_stage(f,initial,np.ones(144),Config(q=0),1,np.full(108,old),np.full(108,old),rule,threshold)

    def test_synthetic_lower_pv_and_emergency_tradeoff(self):
        normal=self.synthetic(pv=50);low=self.synthetic(pv=0)
        self.assertFalse(normal['accepted']);self.assertTrue(low['accepted']);self.assertGreater(low['G'].sum(),normal['G'].sum())
        self.assertGreater(low['voi'],0)

    def test_synthetic_high_pv_no_refund_vs_refund(self):
        a=self.synthetic(pv=200);b=self.synthetic(pv=200,rule='B')
        self.assertFalse(a['accepted']);self.assertTrue(b['accepted']);self.assertLess(b['G'].sum(),a['G'].sum())

    def test_synthetic_small_value_high_threshold(self):
        s=self.synthetic(pv=49.999,threshold=1e6)
        self.assertFalse(s['accepted']);np.testing.assert_array_equal(s['G'],np.full(108,50.))

    def test_synthetic_battery_boundaries(self):
        low=self.synthetic(initial=1200);high=self.synthetic(initial=10800)
        self.assertGreaterEqual(low['G'].sum(),high['G'].sum()-1e-6)
        excess=self.synthetic(pv=300,initial=10800)
        self.assertFalse(excess['accepted'])

    def test_lp_milp_adjustment_equivalence(self):
        f=self.f.feature(80,1);net=f.prediction[:108,0]-f.prediction[:108,1]
        for rule in ('A','B'):
            a=adjustment_plan(net,self.p[36:],6000,np.full(108,400.),terminal_value(self.p),rule)
            b=adjustment_plan(net,self.p[36:],6000,np.full(108,400.),terminal_value(self.p),rule,integer=True)
            self.assertAlmostEqual(a['objective'],b['objective'],places=5)
            self.assertFalse(np.any((a['C']>1e-6)&(a['D']>1e-6)))

    def test_path_execution_matches_q2(self):
        g=np.full(144,400.);h,e=simulate_paths(g,self.a[60:63],6000)
        for j in range(3):
            x=execute_day(g,self.a[60+j],6000)
            np.testing.assert_allclose(h[j],x['H']);self.assertAlmostEqual(e[j],x['E'][-1])

    def test_no_update_identity_and_voi_gate(self):
        s,events=run_day(self.f,60,6000,self.p,Config(),updates=())
        for k in (1,2,3):np.testing.assert_array_equal(s['versions'][k],s['versions'][0])
        self.assertEqual(s['U'].sum()+s['R'].sum(),0)
        s,events=run_day(self.f,60,6000,self.p,Config(scenario=True))
        for ev in events:
            self.assertAlmostEqual(ev['voi'],ev['keep_score']-ev['update_score'])
            if ev['accepted']:self.assertGreater(ev['voi'],1e-6)
            else:self.assertEqual(ev['increase_kwh']+ev['decrease_kwh'],0)

    def test_protected(self):self.assertEqual(check_protected(),0)

    def test_month_selection_future_counterfactual(self):
        from solve_q3 import select_config
        actual=self.a.copy();actual[31:,:,0]*=3;records=dict(self.r)
        for dt,rec in records.copy().items():
            if dt>=datetime(2025,2,1):records[dt]=replace(rec,power=tuple(np.array(rec.power)*.1))
        candidates=(Config(q=0),Config(q=.65),Config(horizon='R48'))
        x,rx=select_config(self.f,31,[6000.]*32,self.p,candidates)
        y,ry=select_config(FeatureFactory(actual,records),31,[6000.]*32,self.p,candidates)
        self.assertEqual(x,y);self.assertEqual(rx,ry)

    def test_scenario_execution_no_future_path_information(self):
        g=np.full(144,400.);paths=self.a[60:62].copy();changed=paths.copy();changed[:,72:,0]*=3
        x,_=simulate_paths(g,paths,6000);y,_=simulate_paths(g,changed,6000)
        np.testing.assert_array_equal(x[:,:72],y[:,:72])
        _,e1=simulate_paths(g[:72],paths[:,:72],6000);_,e2=simulate_paths(g[:72],changed[:,:72],6000)
        np.testing.assert_array_equal(e1,e2)

    def test_T24_18_nextday_forecast_not_visible_to_objective(self):
        day=80;ff=self.f.feature(day,3);pred=ff.prediction.copy();pred[36:,1]*=20
        mutated=replace(ff,prediction=pred)
        old=np.full(36,400.)
        x=choose_stage(ff,6000,self.p,Config(),3,old,old)
        y=choose_stage(mutated,6000,self.p,Config(),3,old,old)
        np.testing.assert_array_equal(x['G'],y['G']);self.assertEqual(x['voi'],y['voi'])

class OutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dates,cls.a,cls.p=read_year();cls.s=dict(np.load(OUT/'experiments/formal_arrays.npz'))

    def test_all_physics_versions_costs(self):
        s=self.s
        for d in range(365):
            validate_execution(self.a[d],{k:s[k][d] for k in ('G','C','D','H','W','E')})
            np.testing.assert_allclose(s['cash'][d],s['versions'][d,0]*self.p+s['up_cost'][d]+s['penalty'][d]-s['refund'][d]+5*self.p*s['H'][d])
            for k in (1,2,3):
                np.testing.assert_array_equal(s['versions'][d,k,:k*36],s['versions'][d,k-1,:k*36])
                np.testing.assert_allclose(s['versions'][d,k],s['versions'][d,k-1]+s['U'][d,k-1]-s['R'][d,k-1],atol=1e-7)
        np.testing.assert_array_equal(s['E'][1:,0],s['E'][:-1,-1])
        np.testing.assert_array_equal(s['G'],s['versions'][:,3])
        self.assertAlmostEqual(float(s['R'][31:].sum()),0,places=5)

    def test_excel_all_values(self):
        s=self.s;w=load_workbook(OUT/'result3.xlsx',read_only=True,data_only=True)
        self.assertEqual(w.sheetnames,['计划购电量','调整购电量','充放电量','紧急购电量'])
        for idx in (0,1):
            rows=list(w.worksheets[idx].values);self.assertEqual(len(rows),335);self.assertEqual(list(rows[0][1:145]),PERIODS)
            for d,row in enumerate(rows[1:],31):
                self.assertEqual(row[0].date(),self.dates[d]);g=s['versions'][d,0] if idx==0 else s['G'][d]
                np.testing.assert_allclose(row[1:145],g,atol=1e-8,rtol=0)
                self.assertAlmostEqual(row[145],g.sum(),places=6)
                self.assertAlmostEqual(row[146],float(g@self.p) if idx==0 else float(s['cash'][d].sum()),places=6)
        rows=list(w.worksheets[2].values)[1:];self.assertEqual(len(rows),334*6)
        for d in range(31,365):
            block=rows[(d-31)*6:(d-30)*6];self.assertEqual(block[0][0].date(),self.dates[d])
            for k,row in enumerate(block):
                self.assertEqual(row[1],f'{4*k}:00-{4*k+4}:00')
                self.assertAlmostEqual(row[2],s['C'][d,k*24:(k+1)*24].sum(),places=6)
                self.assertAlmostEqual(row[3],s['D'][d,k*24:(k+1)*24].sum(),places=6)
            self.assertAlmostEqual(block[0][5],s['E'][d,0],places=6);self.assertAlmostEqual(block[1][5],s['E'][d,-1],places=6)
        events=merge_emergency(self.dates[31:],s['H'][31:]);got=[];date_=None
        for row in list(w.worksheets[3].values)[1:]:
            if row[0]:date_=row[0].date().isoformat()
            if row[1]!='无':got.append((date_,row[1],row[2]))
        self.assertEqual(len(events),len(got))
        for x,y in zip(events,got):
            self.assertEqual((x['date'],x['period']),y[:2]);self.assertAlmostEqual(x['kwh'],y[2],places=6)
        w.close()

    def test_monthly_selection_only_past(self):
        with (OUT/'experiments/monthly_selection.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
        groups=sorted({int(r['selection_day']) for r in rows});self.assertEqual(len(groups),11)
        for d in groups:
            g=[r for r in rows if int(r['selection_day'])==d];self.assertEqual(len(g),len(CANDIDATES))
            self.assertTrue(all(int(r['end'])==d-1 for r in g))
            selected=[r for r in g if r['selected']=='True'];self.assertEqual(len(selected),1)
            self.assertAlmostEqual(float(selected[0]['score']),min(float(r['score']) for r in g),places=5)

    def test_ledger_reconstructs_versions(self):
        with (OUT/'adjustment_ledger.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
        self.assertEqual(len(rows),334*(108+72+36))
        for row in rows:
            day=(datetime.fromisoformat(row['decision_date']).date()-self.dates[0]).days;k=int(row['stage']);t=int(row['index'])
            self.assertAlmostEqual(float(row['previous_commitment']),self.s['versions'][day,k-1,t])
            self.assertAlmostEqual(float(row['new_commitment']),self.s['versions'][day,k,t])
            u=float(row['increase_U']);r=float(row['decrease_R']);p=self.p[t]
            self.assertAlmostEqual(float(row['adjustment_charge']),1.5*p*u+.5*p*r)

    def test_alternative_settlement_and_physics_all_days(self):
        for rule in ('B','B_net'):
            s=dict(np.load(OUT/f'experiments/settlement_{rule}.npz'))
            np.testing.assert_array_equal(s['E'][1:,0],s['E'][:-1,-1])
            self.assertGreater(s['R'][31:].sum(),0)
            for d in range(31,365):
                validate_execution(self.a[d],{k:s[k][d] for k in ('G','C','D','H','W','E')})
                fees=np.zeros(144)
                for k in (1,2,3):
                    fees+=settlement(s['versions'][d,k-1],s['versions'][d,k],self.p,rule,s['versions'][d,0])['charge']
                np.testing.assert_allclose(s['cash'][d],self.p*s['versions'][d,0]+fees+5*self.p*s['H'][d],atol=1e-7)

    def test_reject_ten_minute_shift_and_missing_intervals(self):
        with (OUT/'dispatch_10min.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))[:144]
        labels=[r['period'] for r in rows]
        self.assertEqual(labels,PERIODS)
        self.assertNotEqual(labels[1:]+labels[:1],PERIODS)
        self.assertNotEqual(labels[1:],PERIODS)

if __name__=='__main__':
    core='--core' in sys.argv
    suite=unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(CoreTests)])
    if not core:suite.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(OutputTests))
    result=unittest.TextTestRunner(verbosity=2).run(suite);sys.exit(not result.wasSuccessful())

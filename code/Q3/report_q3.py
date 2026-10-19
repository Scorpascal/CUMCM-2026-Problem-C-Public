"""从原始附件与本次运算数组重建CSV、官方Excel、图件及论文文字。"""
from pathlib import Path
from copy import copy
from datetime import datetime,timedelta
from math import factorial
import csv,json,sys,os
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from solve_q3 import ROOT,OUT,write_json,write_csv,inputs,check_protected
from microgrid_online import PERIODS,merge_emergency,terminal_value
from q3_forecast import FeatureFactory,risk_net
from q3_model import Config,summarize

WANTED=('2025-03-20','2025-06-21','2025-09-23','2025-12-21')
def read_csv(path):
    with path.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def table(headers,rows):
    return '\n'.join(['|'+'|'.join(headers)+'|','|'+'|'.join(['---']*len(headers))+'|']+['|'+'|'.join(map(str,r))+'|' for r in rows])+'\n'
def num(x):return f'{float(x):,.4f}'

def workbook(dates,p,s):
    wb=load_workbook(ROOT/'附件/附件5/result3.xlsx');original=None
    for i,name in enumerate(('计划购电量','调整购电量')):
        ws=wb[name]
        if original is None:original=[ws.cell(1,c).value for c in range(2,146)]
        for t,label in enumerate(PERIODS,2):ws.cell(1,t,label)
        for d in range(31,365):
            row=d-29;g=s['versions'][d,0] if i==0 else s['G'][d]
            ws.cell(row,1,datetime.combine(dates[d],datetime.min.time())).number_format='yyyy-mm-dd'
            for t in range(144):ws.cell(row,t+2,float(g[t])).number_format='#,##0.0000'
            ws.cell(row,146,float(g.sum())).number_format='#,##0.0000'
            ws.cell(row,147,float(p@g) if i==0 else float(s['cash'][d].sum())).number_format='#,##0.0000'
        ws.freeze_panes='B2';ws.auto_filter.ref='A1:EQ335';ws.row_dimensions[1].height=27
        for c in range(1,148):ws.column_dimensions[get_column_letter(c)].width=18 if c in (1,146,147) else 15
    ws=wb['充放电量'];styles=[[copy(ws.cell(r,c)._style) for c in range(1,7)] for r in range(2,8)]
    for merged in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged))
    ws.delete_rows(2,ws.max_row-1)
    for d in range(31,365):
        for k in range(6):
            r=2+(d-31)*6+k
            for c in range(1,7):ws.cell(r,c)._style=copy(styles[k][c-1])
            if k==0:ws.cell(r,1,datetime.combine(dates[d],datetime.min.time())).number_format='yyyy-mm-dd'
            ws.cell(r,2,f'{k*4}:00-{k*4+4}:00')
            for c,key in ((3,'C'),(4,'D')):ws.cell(r,c,float(s[key][d,k*24:(k+1)*24].sum())).number_format='#,##0.0000'
            if k<2:
                ws.cell(r,5,'0:00' if k==0 else '24:00');ws.cell(r,6,float(s['E'][d,0 if k==0 else -1])).number_format='#,##0.0000'
    for c in range(1,7):ws.column_dimensions[get_column_letter(c)].width=19
    ws.freeze_panes='C2'
    ws=wb['紧急购电量'];styles=[copy(ws.cell(2,c)._style) for c in range(1,4)]
    for merged in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged))
    ws.delete_rows(2,ws.max_row-1);events=merge_emergency(dates[31:],s['H'][31:]);r=2
    for d in range(31,365):
        ev=[e for e in events if e['date']==dates[d].isoformat()] or [dict(period='无',kwh=0.)]
        for j,e in enumerate(ev):
            for c in range(1,4):ws.cell(r,c)._style=copy(styles[c-1])
            if j==0:ws.cell(r,1,datetime.combine(dates[d],datetime.min.time())).number_format='yyyy-mm-dd'
            ws.cell(r,2,e['period']);ws.cell(r,3,float(e['kwh'])).number_format='#,##0.0000';r+=1
    for c in range(1,4):ws.column_dimensions[get_column_letter(c)].width=23
    ws.freeze_panes='B2'
    for ws in wb:
        ws.sheet_view.zoomScale=85;ws.print_title_rows='1:1'
    wb.save(OUT/'result3.xlsx');wb.close()
    write_csv(OUT/'time_mapping.csv',[dict(index=t,original_template=original[t],period=PERIODS[t],input_end=PERIODS[t].split('-')[1],column=t+2) for t in range(144)])
    write_csv(OUT/'emergency_purchase_merged.csv',events)
    return events

def data_exports(dates,a,p,s):
    daily=[];dispatch=[];ledger=[]
    for d,date_ in enumerate(dates):
        row=dict(date=date_.isoformat(),config=s['config'][d],initial_kwh=s['E'][d,0],terminal_kwh=s['E'][d,-1],
          plan_kwh=s['versions'][d,0].sum(),effective_kwh=s['G'][d].sum(),emergency_kwh=s['H'][d].sum(),
          increase_kwh=s['U'][d].sum(),decrease_kwh=s['R'][d].sum(),emergency_periods=int((s['H'][d]>1e-7).sum()),
          adjustment_events=int(np.any(s['U'][d]+s['R'][d]>1e-7,axis=1).sum()),load_kwh=a[d,:,0].sum(),pv_kwh=a[d,:,1].sum(),
          pv_curtail_lower_kwh=np.maximum(0,s['W'][d]-s['G'][d]).sum(),pv_curtail_upper_kwh=np.minimum(s['W'][d],a[d,:,1]).sum())
        row.update({k+'_yuan':s[k][d].sum() for k in ('plan_cost','up_cost','penalty','refund','emergency_cost','cash')});daily.append(row)
        if d<31:continue
        local=[]
        for t in range(144):
            rr=dict(date=date_.isoformat(),index=t,period=PERIODS[t],price_yuan_kwh=p[t],load_kw=a[d,t,0]*6,pv_kw=a[d,t,1]*6,
                G0=s['versions'][d,0,t],G6=s['versions'][d,1,t],G12=s['versions'][d,2,t],G18=s['versions'][d,3,t],G_effective=s['G'][d,t],
                C_real=s['C'][d,t],D_real=s['D'][d,t],H_real=s['H'][d,t],W_real=s['W'][d,t],E_start=s['E'][d,t],E_end=s['E'][d,t+1],
                pv_curtail_lower=max(0,s['W'][d,t]-s['G'][d,t]),pv_curtail_upper=min(s['W'][d,t],a[d,t,1]))
            for k in range(4):rr[f'PV_forecast_{k*6}_kw']=s['prediction'][d,k,t,1]*6 if t>=k*36 else ''
            for k in range(3):rr[f'U{6*(k+1)}']=s['U'][d,k,t];rr[f'R{6*(k+1)}']=s['R'][d,k,t]
            rr.update({key+'_yuan':s[key][d,t] for key in ('plan_cost','up_cost','penalty','refund','emergency_cost','cash')})
            dispatch.append(rr);local.append(rr)
        for k in (1,2,3):
            for t in range(k*36,144):
                u=s['U'][d,k-1,t];r=s['R'][d,k-1,t]
                ledger.append(dict(decision_date=date_.isoformat(),decision_time=f'{k*6}:00',stage=k,delivery_date=date_.isoformat(),delivery_interval=PERIODS[t],index=t,
                    previous_commitment=s['versions'][d,k-1,t],new_commitment=s['versions'][d,k,t],increase_U=u,decrease_R=r,normal_price=p[t],
                    increase_charge=1.5*p[t]*u,decrease_penalty=.5*p[t]*r,refund=0.,adjustment_charge=1.5*p[t]*u+.5*p[t]*r,settlement_rule='A_incremental',forecast_version=date_.isoformat()+f'T{k*6:02d}:00'))
        if date_.isoformat() in WANTED:
            write_csv(OUT/'selected_days'/f'{date_}_dispatch.csv',local)
            write_csv(OUT/'selected_days'/f'{date_}_table1.csv',[dict(period=PERIODS[t],G0=s['versions'][d,0,t],effective=s['G'][d,t],emergency=s['H'][d,t]) for t in (60,72,84,96,108,120)])
            write_csv(OUT/'selected_days'/f'{date_}_table2.csv',[dict(period=f'{4*j}:00-{4*j+4}:00',charge=s['C'][d,24*j:24*j+24].sum(),discharge=s['D'][d,24*j:24*j+24].sum()) for j in range(6)])
    write_csv(OUT/'daily_summary.csv',daily[31:]);write_csv(OUT/'warmup_january.csv',daily[:31]);write_csv(OUT/'dispatch_10min.csv',dispatch)
    write_csv(OUT/'adjustment_ledger.csv',ledger)
    write_csv(OUT/'emergency_purchase_10min.csv',[dict(date=r['date'],index=r['index'],period=r['period'],H_real=r['H_real']) for r in dispatch])
    write_csv(OUT/'selected_days/summary.csv',[r for r in daily if r['date'] in WANTED])
    return daily

def forecast_diagnostics(dates,a,p,records,s):
    f=FeatureFactory(a,records);rows=[];dayrows=[];matched=[];flat=a.reshape(-1,2)
    # 固定方法全年误差是事后诊断，正式方法由此前21日经济选择。
    for method in ('constant','linear','pchip'):
        for k in range(4):
            errors=[];targets=[];energies=[];peaks=[]
            for d in range(31,365):
                start=d*144+k*36;n=min(144,len(flat)-start)
                pred=f.raw(d,k,method,0.,k)[:n,1]*6;real=flat[start:start+n,1]*6
                err=pred-real;errors.append(err);targets.append((np.arange(n)+k*36)%144)
                energies.append(err.sum()/6);peaks.append(pred.max()-real.max())
                dayrows.append(dict(date=str(dates[d]),method=method,issue_hour=k*6,mae_kw=np.abs(err).mean(),rmse_kw=np.sqrt((err**2).mean()),energy_bias_kwh=err.sum()/6,peak_bias_kw=pred.max()-real.max()))
            error=np.concatenate(errors);clock=np.concatenate(targets)
            dawn=((clock>=30)&(clock<54))|((clock>=96)&(clock<120))
            rows.append(dict(method=method,issue_hour=k*6,lead_hour='all',n=len(error),mae_kw=np.abs(error).mean(),rmse_kw=np.sqrt((error**2).mean()),bias_kw=error.mean(),mean_energy_bias_kwh=np.mean(energies),dawn_dusk_mae_kw=np.abs(error[dawn]).mean(),mean_peak_bias_kw=np.mean(peaks)))
            for h in range(24):
                e=np.concatenate([er[h*6:(h+1)*6] for er in errors if len(er)>h*6])
                rows.append(dict(method=method,issue_hour=k*6,lead_hour=h+1,n=len(e),mae_kw=np.abs(e).mean(),rmse_kw=np.sqrt((e**2).mean()),bias_kw=e.mean(),mean_energy_bias_kwh='',dawn_dusk_mae_kw='',mean_peak_bias_kw=''))
    for d in range(31,365):
        for k in (1,2,3):
            t=k*36;old=f.raw(d,k,'linear',0.,k-1)[:144-t,1]*6;new=f.raw(d,k,'linear',0.,k)[:144-t,1]*6;real=a[d,t:,1]*6
            matched.append(dict(date=str(dates[d]),issue_hour=k*6,previous_issue_hour=(k-1)*6,change_kwh=np.abs(new-old).sum()/6,
              old_mae_kw=np.abs(old-real).mean(),new_mae_kw=np.abs(new-real).mean(),mae_improvement_kw=np.abs(old-real).mean()-np.abs(new-real).mean()))
    calibration=[]
    calendar=[Config(**c) for c in json.loads((OUT/'experiments/calendar.json').read_text())]
    for k in range(4):
        for obj in (0,1):
            err=[];coverage=[]
            for d in range(31,365):
                c=calendar[d];ff=f.feature(d,k,c.interp,c.gain,c.bias);n=144-k*36
                er=(ff.prediction[:n,obj]-a[d,k*36:,obj])*6;err.extend(er)
                _,m=risk_net(ff,c.q,n);actualnet=a[d,k*36:,0]-a[d,k*36:,1]
                coverage.extend(actualnet<=ff.prediction[:n,0]-ff.prediction[:n,1]+m+1e-7)
            er=np.array(err);calibration.append(dict(object='load' if obj==0 else 'pv',issue_hour=k*6,mae_kw=np.abs(er).mean(),rmse_kw=np.sqrt((er**2).mean()),bias_kw=er.mean(),net_margin_coverage=np.mean(coverage)))
    write_csv(OUT/'forecast_metrics.csv',rows);write_csv(OUT/'experiments/forecast_daily.csv',dayrows);write_csv(OUT/'experiments/forecast_matched_updates.csv',matched);write_csv(OUT/'experiments/calibration.csv',calibration)
    return rows,matched,calibration

def information_value():
    ab=read_csv(OUT/'experiments/information_ablation.csv');mapping={int(r['mask']):r for r in ab if int(r['mask'])>=0};events=read_csv(OUT/'voi_events.csv');rows=[]
    for k in (1,2,3):
        bit=1<<(k-1);values={key:0. for key in ('total_cost_yuan','terminal_adjusted_cost_yuan')}
        for mask in range(8):
            if mask&bit:continue
            count=mask.bit_count();weight=factorial(count)*factorial(2-count)/6
            for key in values:values[key]+=weight*(float(mapping[mask][key])-float(mapping[mask|bit][key]))
        ev=[r for r in events if int(r['stage'])==k];accepted=[r for r in ev if r['accepted']=='True']
        rows.append(dict(issue_hour=k*6,shapley_cash_yuan=values['total_cost_yuan'],shapley_adjusted_yuan=values['terminal_adjusted_cost_yuan'],
          leave_one_out_cash_yuan=float(mapping[7^bit]['total_cost_yuan'])-float(mapping[7]['total_cost_yuan']),
          accepted=len(accepted),opportunities=len(ev),acceptance_rate=len(accepted)/len(ev),voi_positive_rate=np.mean([float(r['voi'])>1e-6 for r in ev]),
          mean_estimated_voi=np.mean([float(r['voi']) for r in ev]),sum_local_realized_value=sum(float(r['realized_value']) for r in ev),
          increase_kwh=sum(float(r['increase_kwh']) for r in ev),decrease_kwh=sum(float(r['decrease_kwh']) for r in ev),
          mean_adjustment_kwh=np.mean([float(r['increase_kwh'])+float(r['decrease_kwh']) for r in ev]),
          mean_solve_seconds=np.mean([float(r['seconds']) for r in ev]),max_solve_seconds=max(float(r['seconds']) for r in ev)))
    write_csv(OUT/'experiments/stage_value.csv',rows)
    return ab,rows

def figures(dates,a,p,s,daily,metrics,stage):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':130})
    directory=OUT/'figures';directory.mkdir(exist_ok=True)
    def save(fig,name):
        fig.savefig(directory/(name+'.png'),dpi=160,bbox_inches='tight');fig.savefig(directory/(name+'.svg'),bbox_inches='tight');plt.close(fig)
    x=(np.arange(144)+.5)/6
    for date_ in WANTED:
        d=next(i for i,dt in enumerate(dates) if str(dt)==date_)
        fig,ax=plt.subplots(6,1,figsize=(12,16),sharex=True,layout='constrained')
        ax[0].plot(x,a[d,:,0]*6,label='Actual load',color='#334155');ax[0].plot(x,a[d,:,1]*6,label='Actual PV',color='#f59e0b',lw=2)
        for k in range(4):ax[0].plot(x,s['prediction'][d,k,:,1]*6,label=f'PV issue {k*6:02d}:00',lw=1,alpha=.8)
        ax[0].set_ylabel('Power (kW)');ax[0].legend(ncol=3,fontsize=8);ax[0].set_title(date_+' | Forecasts, commitments and actual operation',loc='left')
        for k in range(4):ax[1].step(x,s['versions'][d,k],where='mid',label=f'G{k*6}',alpha=.8)
        ax[1].plot(x,a[d,:,0]-a[d,:,1],color='#64748b',lw=.7,label='Actual net load');ax[1].set_ylabel('kWh / 10 min');ax[1].legend(ncol=5,fontsize=8)
        for k in range(3):ax[2].step(x,s['U'][d,k]-s['R'][d,k],label=f'Update {6*(k+1)}h',where='mid')
        ax[2].axhline(0,color='#94a3b8',lw=.5);ax[2].set_ylabel('Adjustment (kWh)');ax[2].legend(ncol=3,fontsize=8)
        ax[3].fill_between(x,0,s['C'][d]*6,label='Charge',alpha=.5,color='#10b981');ax[3].fill_between(x,0,-s['D'][d]*6,label='Discharge',alpha=.5,color='#3b82f6')
        ax[3].step(x,s['H'][d]*6,label='Emergency',color='#dc2626',where='mid');ax[3].set_ylabel('Power (kW)');ax[3].legend(ncol=3,fontsize=8)
        ax[4].plot(np.arange(145)/6,s['E'][d],color='#7c3aed');ax[4].axhline(1200,ls=':',color='#94a3b8');ax[4].axhline(10800,ls=':',color='#94a3b8');ax[4].set_ylabel('Stored energy (kWh)')
        ax[5].plot(x,np.cumsum(s['cash'][d]),color='#0f766e',label='Cash allocated by delivery');ay=ax[5].twinx();ay.plot(x,p,color='#ef4444',alpha=.6);ay.set_ylabel('Price (yuan/kWh)');ax[5].set_ylabel('Cumulative yuan');ax[5].set_xlabel('Hour of delivery (costs allocated, not payment timestamp)')
        for axes in ax:
            axes.set_xlim(0,24);axes.set_xticks(np.arange(0,25,2));axes.grid(alpha=.15)
            for tick in (6,12,18):axes.axvline(tick,color='#94a3b8',ls='--',lw=.6)
        save(fig,date_)
    fig,ax=plt.subplots(1,2,figsize=(12,4.4),layout='constrained')
    for k in (0,6,12,18):
        rr=[r for r in metrics if r['method']=='linear' and r['issue_hour']==k and r['lead_hour']!='all']
        ax[0].plot([r['lead_hour'] for r in rr],[r['mae_kw'] for r in rr],label=f'Issue {k}h')
    ax[0].set(xlabel='Lead time (h)',ylabel='PV MAE (kW)',title='Raw linear forecasts; comparable lead groups');ax[0].legend()
    ax[1].bar([str(r['issue_hour']) for r in stage],[r['shapley_cash_yuan']/1000 for r in stage],color=['#0f766e','#3b82f6','#f59e0b'])
    ax[1].set(xlabel='Additional forecast issue hour',ylabel='Cash saving (thousand yuan)',title='Shapley allocation across all eight subsets');ax[1].axhline(0,color='#64748b',lw=.5);save(fig,'forecast_and_value')
    fig,ax=plt.subplots(figsize=(11,4.5),layout='constrained');months=range(2,13);bottom=np.zeros(11)
    for key,label,color in [('plan_cost_yuan','Original plan','#334155'),('up_cost_yuan','Increase','#38bdf8'),('emergency_cost_yuan','Emergency','#ef4444')]:
        vals=np.array([sum(float(r[key]) for r in daily if r['date'][5:7]==f'{m:02d}') for m in months])/1e6
        ax.bar(list(months),vals,bottom=bottom,label=label,color=color);bottom+=vals
    ax.set(xlabel='Month',ylabel='Million yuan',title='Actual cash cost by month');ax.legend(ncol=3);save(fig,'monthly_costs')

def documents(dates,a,p,s,daily,merged,metrics,matched,calibration,ablation,stage):
    summary=summarize(s,p);q2=json.loads((ROOT/'Q2/summary.json').read_text())
    subsets={int(r['mask']):r for r in ablation};only0=subsets[0]
    candidates=read_csv(OUT/'experiments/candidate_summary.csv');sensitive=read_csv(OUT/'experiments/sensitivities.csv')
    choices=[r for r in read_csv(OUT/'experiments/monthly_selection.csv') if r['selected']=='True'];events=read_csv(OUT/'voi_events.csv')
    supplemental=read_csv(OUT/'experiments/supplemental.csv');supp={r['policy']:r for r in supplemental}
    independent0=supp['zero_only_independent_past_selection'];independent_gain=float(independent0['total_cost_yuan'])-summary['total_cost_yuan']
    cash=summary['total_cost_yuan'];gain_q2=q2['total_cost_yuan']-cash;gain0=float(only0['total_cost_yuan'])-cash
    policy_counts={c:int(np.sum(s['config'][31:]==c)) for c in np.unique(s['config'][31:])}
    state_effect=float(subsets[-1]['total_cost_yuan'])-cash
    text=f'''# 问题3 基于已发布光伏预报的多阶段购电调整

## 1 问题分析与结论

问题2在0时冻结全天普通购电，问题3增加0/6/12/18时官方光伏预报，并允许在日内按规定价格改变未交付承诺。新信息应在节省紧急购电与支付调整费用之间权衡。本研究保留已验收Q2的实时储能反馈，采用预注册候选的月度历史经济选择，在每次预报发布时形成或调整未来承诺。

主解释为原计划费用不退、调增按1.5倍、调减另付0.5倍，按上一有效承诺逐事件结算。2025-02-01至12-31共334天实际费用为{num(cash)}元，相对冻结Q2现金节省{num(gain_q2)}元（{gain_q2/q2['total_cost_yuan']:.2%}）。紧急购电{num(summary['emergency_kwh'])} kWh，共{summary['emergency_periods']}个十分钟时段；Q2为146312.1414 kWh、923时段。因此费用及缺口电量改善不等于所有可靠性指标同时改善。

独立历史选参的仅0时预报对照费用{num(independent0['total_cost_yuan'])}元，完整日内更新比它节省{num(independent_gain)}元；它相对Q2费用增加{num(float(independent0['total_cost_yuan'])-q2['total_cost_yuan'])}元，不能声称官方0时预报单独必然优于历史预测。另有共享正式配置日历的仅0时对照{num(only0['total_cost_yuan'])}元，用于隔离同一风险配置下的调整权限。以下结论只适用于声明的结算、执行器、候选集和历史数据。正式方法是在线选择器，不能把事后全年最低固定候选替换成2月份已提交计划。

## 2 事实 假设与信息集

原题规定每日四次未来24小时整点PV预报、调整价与紧急价、储能上下界和结果日期。沿用Q2建模解释：10分钟记录为区间结束点、功率/6近似电量、充放电效率各0.9、实时反馈、不售电、不削负荷、无额外并网容量。实时反馈使用当前时段观测的聚合近似，不是时段开始前就知道其完整平均值。Q3没有重复Q1每日首末相等要求，全年SOC连续。

F(d,k)包括阶段开始前全部已实现记录、真实当前SOC、已发布预报、固定价格、日历、已完成历史误差和月初已确定配置；不包括尚未发布预报或未来实际曲线。0时只能读取当日0时预报；6时增加0–6实际与6时预报；12/18同理。实际未来只进入逐时执行环境和完成后的评价，不进入choose_stage接口。

1月1日初值6000 kWh。因缺历史负载，采用Q2同类保守启动日：G=0、电池静置、缺口补电；1月2日起固定Q3热启动，1月费用另存。所有Q3实验共享这条1月轨迹。2月1日SOC为{num(summary['initial_kwh'])} kWh、年底为{num(summary['terminal_kwh'])} kWh；Q2相应为10800和3453.7471 kWh，不能隐去这一边界差异。

## 3 附件3结构及整点映射

附件3含365×4=1460次发行、每次24个点，显式键为发行日期、发行时间、目标时间和提前小时。0时目标为1时至次日0时；6/12/18时依次平移，年末部分目标落在2026年，未提供的实际值不参与误差评价。原始35040个点保存在forecast_versions.csv。

比较向后保持常值、线性和PCHIP。发行时最后已知PV为左端锚点，首次1月1日0时取0；原始整点值必须保持，插值为本文选择。区间结束点采样除以6近似电量，不声称小时功率预报已给出小时总能量。实际月度选择及误差见下表；全年误差仅诊断，正式插值由过去21日经济回放选择。

'''
    text+=table(['方法','发行时刻','PV MAE kW','RMSE kW','预报减实际bias kW','24h平均能量偏差kWh','晨昏MAE kW'],[[r['method'],r['issue_hour'],num(r['mae_kw']),num(r['rmse_kw']),num(r['bias_kw']),num(r['mean_energy_bias_kwh']),num(r['dawn_dusk_mae_kw'])] for r in metrics if r['lead_hour']=='all'])
    text+=r'''
最后一日的部分24h目标不可观测，因此能量统计使用可见目标子集，完整日诊断另见forecast_daily.csv；晨昏定义为5–9及16–20时。首小时常值使用发行锚点向后保持，可能有晨昏滞后，不能把不同插值误差视为同一算法的时间错位。

## 4 负载预测 官方预报校准与风险

负载沿用Q2组合预测：最近7日同刻均值与过去最多4个同星期同刻均值各半，整个未来窗口固定同一历史截断点。候选日内修正使用最近最多6小时已观测负载误差均值，以6小时衰减时间指数递减；增益0/.5/1由历史验证选择。

官方PV优先覆盖其24h有效范围，之外由Q2历史PV预测补齐。偏差校正候选用过去对应发行阶段、相对时距的误差均值；零样本不校正，不按全年数据拟合。误差档案仅收录已完整实现的48h负载/PV联合路径，最多30条；这比只等待目标点实现更保守，牺牲最新尾部信息以保证实现简洁。分位裕度为

$$m_{d,k,t}=\max\{0,Q_q(e^L-e^{PV})\},\qquad \widetilde N=\widehat L-\widehat P+m. \tag{1}$$

q=0是无裕度点预测的标记，不表示0分位数。固定加性PV bias若对应净误差分位仍为正，会被同量分位平移抵消；因此该组合下费用一致有代数原因，不能据此断言所有PV校准无效。经验覆盖率不具有独立同分布或conformal严格保证。

'''
    text+=table(['对象','阶段h','正式预测MAE kW','RMSE kW','净负荷裕度经验覆盖'],[[r['object'],r['issue_hour'],num(r['mae_kw']),num(r['rmse_kw']),f"{r['net_margin_coverage']:.2%}"] for r in calibration])
    text+=r'''
## 5 数学模型

日期d、阶段k、交付时段t；G0为0时原计划，Gᵏ为当前承诺，U/R为本轮调增/调减，C/D/H/W/E为真实充电、放电、紧急购电、未利用能量和储电量。所有流量kWh，p元/kWh。

$$G_t^k=G_t^{k-1}+U_t^k-R_t^k,\quad U_t^k,R_t^k\ge0;\quad G_t^k=G_t^{k-1},\ t<\tau_k. \tag{2}$$

$$G_t+P_t+D_t+H_t=L_t+C_t+W_t,\qquad E_{t+1}=E_t+0.9C_t-D_t/0.9. \tag{3}$$

$$1200\le E_t\le10800,\quad 0\le C_t\le(5000/6)z_t,\quad0\le D_t\le(5000/6)(1-z_t),\quad z_t\in\{0,1\}. \tag{4}$$

G/H/W非负，实际H只补剩余缺口。预测提案层用G+D−C−W=风险修正净负荷，且不预先把实际H当作普通购电。0时提案目标为ΣpG−vE末；调整提案主A目标为当前日Σ(1.5pU+0.5pR)+未来日ΣpG−vE末。G0在F(d,0)可测，Gᵏ在F(d,k)可测，跨日E[d+1,0]=E[d,144]。

终端代理v=当日固定电价最低24个点均值/0.9=0.4676666667元/内部kWh。T24至自然日末，R24/R48分别从发行时刻向后24/48h；均采用同一v，不每日闭合。未来日购电仅前瞻辅助，次日0时重新形成正式G0，不能把今晚前瞻重复记账。

LP计划允许W作为总未利用能量，若同充同放，令δ=min(C,D/(ηcηd))，扣减C的δ、D的ηcηdδ并增加W的(1−ηcηd)δ，G、E与费用不变，恢复互斥。代表性调整LP与显式MILP目标对照，实际反馈天然互斥。该结论针对允许未利用电量的本模型，不应推广到其他弃电边界。

## 6 实时执行与调整账本

每时段令r=G+PV−L。r≥0时C=min(r,5000/6,(10800−E)/0.9)，D=H=0，W=r−C；r<0时D=min(−r,5000/6,(E−1200)0.9)，C=W=0，H=−r−D。此为继承Q2的反馈器，不是全局最优储能控制，不向电池馈入未来场景。紧急购电只满足负载缺口，不能用于主动充电。

主账单为

$$C_A=\sum_t p_tG_t^0+\sum_{k=1}^3\sum_{t\ge\tau_k}(1.5p_tU_t^k+0.5p_tR_t^k)+\sum_t5p_tH_t. \tag{5}$$

每笔调整相对上一有效承诺，不再对完整新G收一次计划费。调减后电量不再参加物理平衡，即使原费用已沉没。主A下付费调减被保留更多承诺并弃用的策略支配，本次R=0与这一逻辑一致。

替代B退回pR再收0.5pR，调减净费−0.5pR；B_net相对原始G0净偏差重新计账并只收新旧账面差，允许取消未交付早先调整，但保留每笔冲销。原计划100依次80、110、90，p=1时A/B/B_net非紧急总费分别165/125/95。题面未唯一消除退费及基准歧义，主口径不是官方认证。

## 7 更新价值与有限场景方法

在当前相同预测路径、SOC和费用基础下，同时评价KEEP与UPDATE。对给定G逐点调用同型反馈器，计算

$$J=\Delta C_{\rm adj}+\mathbb E_s\big[5\sum pH_s+C_{\rm future}-v(E_{T,s}-E_{\rm now})\big],\quad VOI=J_{\rm keep}-J_{\rm update}.\tag{6}$$

已沉没计划费对两者相同，不重复计入比较。门限1e−6元仅为浮点容差；相等先KEEP。J已包含调整费，不能再从VOI扣一次。跨日前瞻在当前平均路径得到预计午夜E后产生一条共享未来普通购电曲线，所有场景使用同一曲线；不是场景实现后自由重解整个路径。该有限策略近似没有未来发行预报的场景树，不声称多阶段随机最优。

场景候选使用最多8条历史完整联合残差加当前预测，不逐点独立噪声；从q=0/.65/.8/.9提案中选择平均J最小者。当前G跨所有场景共享；实际储能始终逐点反馈，因此没有全知场景电池。当前规模无需K-medoids；未添加CVaR权重，不代表已验证所有CVaR方法无收益。

estimated VOI只使用当前信息。realized_value是决策提交后用真实余日分别冻结KEEP与已采用方案计算的现金差和终端值差，不模拟后续再更新，不能把多个事件的这一重叠诊断值相加当全年预报贡献。

## 8 候选选择与严格滚动验证

16个配置在正式全年排名之前注册，包含点预测、不同风险分位、插值、偏差修正、负载更新、T24/R24/R48、有限场景与取消门控。每月首日用过去21个完整日，所有候选从同一个已知验证起点SOC独立连续回放，评分为现金费−v(E末−E初)，最小者月内固定。正式未来实现仍按顺序揭示。以下为实际选择，并非全年事后最优回填。

    '''
    text+=table(['月首','选择配置','历史窗口','历史评分元'],[[str(dates[int(r['selection_day'])]),r['config'],f"{dates[int(r['begin'])]}—{dates[int(r['end'])]}",num(r['score'])] for r in choices])
    text+='\n固定候选全年诊断（各候选相同Q3热启动，未用于回选过去）：\n\n'
    text+=table(['配置','现金元','末态调整元','紧急kWh','期末E'],[[r['config'],num(r['total_cost_yuan']),num(r['terminal_adjusted_cost_yuan']),num(r['emergency_kwh']),num(r['terminal_kwh'])] for r in candidates])
    text+='\n## 9 全年费用与运行结果\n\n'
    text+=table(['指标','结果'],[[k,num(v)] for k,v in summary.items()])
    lower=sum(r['pv_curtail_lower_kwh'] for r in daily[31:]);upper=sum(r['pv_curtail_upper_kwh'] for r in daily[31:]);pv=a[31:,:,1].sum()
    text+=f'\n未利用能量不等于弃光。按光伏优先分摊，弃光为下界{num(lower)} kWh，消纳率{1-lower/pv:.2%}；不固定来源时弃光可在{num(lower)}至{num(upper)} kWh之间。总费用和供需平衡不依赖该来源分配。\n\n![月度费用](figures/monthly_costs.png)\n'
    text+='\n## 10 后续预报信息消融\n\n'
    text+=table(['信息/策略','现金元','末态调整元','紧急kWh','调整次数','期初E','期末E'],[[r['policy'],num(r['total_cost_yuan']),num(r['terminal_adjusted_cost_yuan']),num(r['emergency_kwh']),r['adjustment_events'],num(r['initial_kwh']),num(r['terminal_kwh'])] for r in ablation])
    text+=f'''\n8个预报子集沿用正式策略的因果配置日历，拥有各自连续SOC，并非分别再调参后的最优子策略。旧预报对照在同样6/12/18时更新SOC、允许同样付费调整，但PV保持0时发行版本。全预报相对旧预报重优化节省{num(state_effect)}元，才更接近同权限下新PV信息的增量贡献；全预报相对全天不调整的{num(gain0)}元同时包含重新观测SOC和合同调整机会。

对8子集用三变量Shapley平均边际节省分配6/12/18的交互影响，三者之和严格等于仅0时费用减全预报费用。leave-one-out为保留另两个时刻时加入该时刻的边际；二者均是事后策略比较，不作为当天VOI。每日现金分配到交付时段，图中累计不是交易付款时间。

'''
    text+=table(['时刻h','Shapley现金节省元','末态调整节省元','移除此时刻增加费用元','采用次数/334','VOI正比例','平均预计VOI元','平均/最大求解秒'],[[r['issue_hour'],num(r['shapley_cash_yuan']),num(r['shapley_adjusted_yuan']),num(r['leave_one_out_cash_yuan']),f"{r['accepted']}/334",f"{r['voi_positive_rate']:.2%}",num(r['mean_estimated_voi']),f"{r['mean_solve_seconds']:.5f}/{r['max_solve_seconds']:.5f}"] for r in stage])
    text+='\n### 同调整权限下的新PV信息与公平基准\n\n'
    text+=table(['补充对照','现金元','紧急kWh','末态调整元'],[[r['policy'],num(r['total_cost_yuan']),num(r['emergency_kwh']),num(r['terminal_adjusted_cost_yuan'])] for r in supplemental])
    text+='\n逐个屏蔽该时刻的新PV但保留当前SOC更新及全部调整权限，其余时刻正常发行。以“屏蔽费−完整费”为新PV的条件现金价值：'
    fresh=[]
    for k in (6,12,18):
        val=float(supp[f'withhold_new_{k}_same_rights']['total_cost_yuan'])-cash;fresh.append((k,val))
        text+=f'{k}时{num(val)}元；'
    text+='新PV边际存在交互，不与前三个Shapley权限值相加。6时新PV在这组条件下现金边际为负但可降低紧急电量，不能声称每份新预报都值得采用，也不能把它的事后负值用于撤销历史正式决策。独立仅0时基准仍只用过去21日选参，属于发现共享低q配置对不可调整基准不公平后增加的诊断，不改变预注册正式策略。\n'
    text+='\n![预报误差与经济价值](figures/forecast_and_value.png)\n\n## 11 18时预报与替代口径\n\n'
    text+=table(['策略','现金元','末态调整元','调增kWh','调减kWh','紧急kWh','末E'],[[r['policy'],num(r['total_cost_yuan']),num(r['terminal_adjusted_cost_yuan']),num(r['increase_kwh']),num(r['decrease_kwh']),num(r['emergency_kwh']),num(r['terminal_kwh'])] for r in sensitive])
    sens={r['policy']:r for r in sensitive}
    for horizon in ('T24','R48'):
        delta=float(sens[horizon+'_without18']['total_cost_yuan'])-float(sens[horizon+'_with18']['total_cost_yuan'])
        text+=f'\n{horizon}固定q=.8对照中，加入18时更新的全年现金节省为{num(delta)}元。'
    nextvalue=float(supp['R48_18_nextday_old_information']['total_cost_yuan'])-float(sens['R48_with18']['total_cost_yuan'])
    text+=f'\n进一步保持18时当日目标、SOC与调整权限不变，仅将次日预报目标替换为12时已知预测/因果补齐；R48的次日新信息现金价值为{num(nextvalue)}元。本数据和该候选下没有观测到额外的次日准备节省，不能照搬报告的“18时次日PV必然有用”。\n'
    text+='''\nT24的固定终端价值不依赖次日PV，因此仅有次日预报变化时无法使T24决策改变；不能给T24虚构跨日天气价值。R48可使用次日PV，以上R48对照同时含18时SOC更新与调整权限的效应，不能仅凭它证明“次日PV本身”导致节省。正式新/旧预报同权限对照和匹配目标误差共同约束解释。B/B_net数值是相同历史配置日历下重新优化和执行，不是只对主轨迹换收费公式，也未对替代口径重新寻找所有最优超参数。

## 12 四个指定日期

表1报告最后实际生效普通购电G，原计划G0并列以展示调整；H另列，不与普通调整混淆。全天费用包括所有费用；四小时段的充放电均为真实执行汇总。日首末SOC可能不同，不能把透支储能产生的表面收益归因于预报。
'''
    selected=[]
    for date_ in WANTED:
        d=next(i for i,dt in enumerate(dates) if str(dt)==date_);row=daily[d];selected.append(row)
        text+=f'\n### {date_}\n\n表1 指定时段购电量（kWh）\n\n'
        text+=table(['时间段','原计划G0','有效G','紧急H'],[[PERIODS[t],num(s['versions'][d,0,t]),num(s['G'][d,t]),num(s['H'][d,t])] for t in (60,72,84,96,108,120)])
        text+=f"\n全天原计划{num(row['plan_kwh'])} kWh，有效普通购电{num(row['effective_kwh'])} kWh，紧急{num(row['emergency_kwh'])} kWh；原计划费{num(row['plan_cost_yuan'])}元、调增费{num(row['up_cost_yuan'])}元、调减费{num(row['penalty_yuan'])}元、紧急费{num(row['emergency_cost_yuan'])}元，总费{num(row['cash_yuan'])}元。\n\n表2 真实储能充放电（kWh）\n\n"
        text+=table(['时间段','充电','放电'],[[f'{4*k}:00-{4*k+4}:00',num(s['C'][d,k*24:(k+1)*24].sum()),num(s['D'][d,k*24:(k+1)*24].sum())] for k in range(6)])
        text+=f"\n0:00储电量{num(s['E'][d,0])} kWh，24:00储电量{num(s['E'][d,-1])} kWh。\n\n"
        ev=[r for r in events if r['date']==date_];mt=[r for r in matched if r['date']==date_]
        text+=table(['更新h','PV曲线变化kWh','匹配目标MAE改善kW','预计VOI元','采用','局部事后价值元'],[[r['issue_hour'],num(r['change_kwh']),num(r['mae_improvement_kw']),num(ev[j]['voi']),ev[j]['accepted'],num(ev[j]['realized_value'])] for j,r in enumerate(mt)])
        accurate_but_keep=[r['issue_hour'] for j,r in enumerate(mt) if r['mae_improvement_kw']>1e-7 and ev[j]['accepted']=='False']
        text+=f'\n本日“匹配目标预报更准但未采用购电调整”的时刻：{accurate_but_keep if accurate_but_keep else "没有出现"}。若没有出现也不预设它必然存在；全年事件可在forecast_matched_updates与voi_events按日期阶段联合复核。\n'
        largest=max(mt,key=lambda r:r['change_kwh']);best=max(mt,key=lambda r:r['mae_improvement_kw']);economic=max(ev,key=lambda r:float(r['realized_value']))
        text+=f"\n本日PV变化量最大在{largest['issue_hour']}时，匹配目标误差改善最大在{best['issue_hour']}时，已采用策略相对冻结KEEP的局部事后价值最大在{economic['decision_time']}。这些排名不要求相同：电池余量、未来缺口、1.5倍增购价与5倍紧急价共同决定是否调整。局部事后价值可为负，不能反过来撤销当时已提交决策。\n\n![{date_}调度](figures/{date_}.png)\n"
    text+='\n### 表3 四指定日紧急购电\n\n'
    evs=[[r for r in merged if r['date']==date_] for date_ in WANTED];tabrows=[]
    for j in range(max(1,max(map(len,evs)))):
        row=[]
        for ev in evs:row.extend([ev[j]['period'],num(ev[j]['kwh'])] if j<len(ev) else ['无' if j==0 else '',0 if j==0 else ''])
        tabrows.append(row)
    text+=table([x for date_ in WANTED for x in (date_+' 时段','kWh')],tabrows)
    text+='''
## 13 验证与结果文件

原型覆盖1/29–2/1、3/19–3/22、6/19–6/22、12/18–12/21四个连续窗口，每窗16候选。核心测试包括整点锚点、功率单位、预报跨日、残差全部实现、修改未来发行版本和未来实际曲线的反事实、过去计划冻结、A/B/B_net手算多交易、缺电增购/富余不调整、SOC边界、LP/MILP对照与继承Q2执行器一致。正式后另全量验证365天物理、334×144计划/调整、充放电和紧急合并，独立从CSV重算费用，并重跑Q1/Q2/Q3全部测试，状态见validation.json与Q3最终验收审计.md。

result3由官方模板填写，计划表144点及全天量是G0，计划费列为原计划pG0；调整表144点及全天量是最终有效G，费用列是最终实际总账单。两表费用不能相加。这一主语义与Q2计划表费用含紧急的既有写法不同，已在口径核验公开，底层费用无混用。内部保留G0/G6/G12/G18、U/R、交易账本和真实执行，任何表均可重建。时间列统一自然日，无循环移动数值；原始模板不改。

## 14 模型评价与局限

该策略的优点是发行时间隔离、冻结承诺与实际执行分层、逐笔账本、历史选参及完整输出核验。它优化的是有限候选类，不是全年随机控制全局最优；贪心储能未优化未来应急机会成本。48h补齐和末态值均为代理，无未来价值精确保证；1月启动、人为残差窗口与候选集设计使用了题目总体结构，因此只称严格时间样本外回放，不称外部独立数据集验证。

最主要WARNING是调减退费/多次调整基准、时间与效率歧义、十分钟即时反馈假设、预报子集共享配置日历、末态SOC不同以及局部VOI与全年价值不可混用。主口径A无调减是结构性经济结果，不表示题目50%条款被遗忘。更低紧急电量伴随更多小缺口时段，需根据实际偏好判断可靠性取舍。当前实验不证明CVaR、复杂MPC或机器学习在所有数据上无用，只是不凭名称增加未验证复杂度。

## 参考资料

[1] C题.pdf，问题2/3、附录1/2与result3合同；附件1/2/3/5。

[2] Q2冻结决策报告、正式论文与最终验收审计，Q2_frozen_v1。

[3] Polimeni et al. Implementation of Different PV Forecast Approaches in a MultiGood MicroGrid: Modeling and Experimental Results. Processes, 2021, 9:323. DOI:10.3390/pr9020323。固定执行框架比较预报、末态校正及精度与经济性区别，见第5/10/15页。

[4] Wen et al. Modeling of photovoltaic power uncertainties for impact analysis on generation scheduling and cost of an urban micro grid. Mathematics and Computers in Simulation, 2021,183:116–128. DOI:10.1016/j.matcom.2020.02.023。本地未校正稿第4–6、9–11页，借鉴误差与风险经济性，不照搬当年未来样本或机组备用。

[5] Panda et al. Methods for Mitigating Uncertainty in Real-Time Operations of a Connected Microgrid. SEGAN,2024,38:101334. DOI:10.1016/j.segan.2024.101334。承诺/实时执行分层。

[6] Herding et al. Stochastic programming approach for optimal day-ahead market bidding curves of a microgrid. Applied Energy,2023,336:120847. DOI:10.1016/j.apenergy.2023.120847。锁定第一阶段及模拟样本外局限，第7页。
'''
    (OUT/'Q3_final.md').write_text(text)
    decision=f'''# Q3最终决策报告

正式策略为16个预注册候选中的月度过去21日经济选择，不使用全年排名回填。实际配置及使用天数：{json.dumps(policy_counts,ensure_ascii=False)}。

0时官方PV进入G0；负载继承Q2组合；日内根据当前SOC、已发布预报和历史残差提出调整，在同一信息下比较KEEP/UPDATE。实际执行完全复用Q2 feedback_step。根目录新增microgrid_forecasts.py及microgrid_adjustment.py供Q4复用，原公共模块/Q1/Q2未修改。

总费{num(cash)}元：原计划{num(summary['plan_cost_yuan'])}，调增{num(summary['up_cost_yuan'])}，调减违约{num(summary['penalty_yuan'])}，退款{num(summary['refund_yuan'])}，紧急{num(summary['emergency_cost_yuan'])}。紧急电量{num(summary['emergency_kwh'])} kWh、{summary['emergency_periods']}时段；调增{num(summary['increase_kwh'])} kWh、调减{num(summary['decrease_kwh'])}，{summary['adjustment_events']}次阶段调整。

独立历史选择Q3仅0时相对Q2现金变化（正为节省）{num(q2['total_cost_yuan']-float(independent0['total_cost_yuan']))}元；完整Q3相对独立仅0时节省{num(independent_gain)}元。共享同一正式配置日历的仅0时消融费用为{num(only0['total_cost_yuan'])}元，完整Q3比它节省{num(gain0)}元；这不是两个均独立调优的对照。完整Q3相对Q2节省{num(gain_q2)}元。不同期初/末SOC应同时查看末态调整费用，不能全部归因于预测。

'''
    decision+=table(['阶段h','Shapley现金价值元','移除此阶段增费元','采用次数'],[[r['issue_hour'],num(r['shapley_cash_yuan']),num(r['leave_one_out_cash_yuan']),r['accepted']] for r in stage])
    decision+='\n在保持调整权限和SOC更新的逐个新PV屏蔽对照中，6/12/18时新预报条件现金价值分别为'+ '、'.join(num(v)+'元' for _,v in fresh)+'。6时该边际略负，不能宣称所有新PV都降低费用。18时次日PV专项R48对照的现金价值为'+num(nextvalue)+'元，未观察到报告预设的跨日收益。\n'
    decision+='''
预报使用不等于每次必须改变G。estimated VOI完全事前；realized value只在提交后计算，且是冻结余日的局部诊断，不能加总为年度因果贡献。是否保留6/12/18参见上表的子集比较与论文18时专节，不以太阳落山直接判断。

原始整点锚点保持，插值、bias、负载更新、时域、场景及门控由月首历史选择决定。R24/R48是候选，未经选择不得宣布它们主用。有限场景不是全知电池的自由两阶段规划；无K-medoids与CVaR，不声称已证明所有复杂扩展无收益。全年固定候选排名、bias代数抵消和窗口结果见Q3_final第8节。

主A不退款，调增1.5p、调减+.5p，针对上一承诺按交付价格记账；B/B_net另重算，结果不能混用。A下没有调减有经济支配理由。result3基于原始模板，计划表为G0，调整表为有效G，充放电与H均来自实际执行；全量回读与底层独立汇总的最终状态以validation.json和Q3最终验收审计.md为准。

四指定日期：

'''
    decision+=table(['日期','原计划费','调增费','紧急费','总费','紧急kWh','首/末E'],[[r['date'],num(r['plan_cost_yuan']),num(r['up_cost_yuan']),num(r['emergency_cost_yuan']),num(r['cash_yuan']),num(r['emergency_kwh']),f"{num(r['initial_kwh'])}/{num(r['terminal_kwh'])}"] for r in selected])
    best=max([ev for ev in events if ev['date'] in WANTED],key=lambda ev:float(ev['realized_value']))
    decision+=f"\n四指定日中，单次余日冻结对照的最大局部更新价值发生在{best['date']} {best['decision_time']}，为{num(best['realized_value'])}元；不把同一天重叠余日价值直接相加当成日因果收益。\n"
    decision+='\nWARNING：结算歧义、效率/时间假设、实时反馈聚合、终端代理与边界不同、同日历消融及局部VOI不可加总，均在论文中披露。紧急电量下降但时段次数增加，不应写成所有可靠性指标改善。\n'
    (OUT/'Q3最终决策报告.md').write_text(decision)
    (OUT/'README.md').write_text('''# 问题3程序与交付

正式文件result3.xlsx，论文Q3_final.md，选择与结算说明见Q3最终决策报告.md及各审计。主A，不退款、逐事件上一承诺结算；B/B_net仅敏感性。本目录是公开衍生副本，包含Q1/Q2计算依赖，未包含私有原始冻结件。

在code目录，先按公开仓库的数据说明自行准备附件1、2、3及附件5中的result1、2、3模板，再使用已安装公开依赖的Python：

```bash
python reproduce.py --through Q3 --check-inputs
python reproduce.py --through Q3
```

入口先计算Q1/Q2，再执行Q3预注册文档、原型、核心测试、正式历史选择、全部实验、交付与数值审计。Q3/solve_q3.py可单独使用--prototype、--compute、--deliver；--compute需要先运行原型，report_q3.py还需要已生成的Q2/summary.json。请在独立公开副本中重算，计算会写入该副本的Q1/Q2/Q3目录。

audit_q3.py独立读取CSV/Excel复算并运行Q1/Q2/Q3测试，日志全部在Q3。公开包不包含原始冻结脚本、清单、交接快照或ZIP；缺省冻结核验明确记录为not_packaged。可用依赖见公开仓库的config/requirements.txt。图表英文以避免中文字体缺失，中文解释在论文。

工作簿由openpyxl导出并由Python测试回读，不需要Node或专有渲染工具。数值审计不等于人工视觉核对或私有原件的完整性证明。

重要输出：dispatch_10min.csv保存G0/G6/G12/G18/有效G、U/R和实际C/D/H/E；adjustment_ledger.csv逐事件账本（含0调整记录）；forecast_versions.csv原始35040预报；voi_events.csv事前VOI与局部事后诊断；experiments保存预注册、月首验证、固定候选诊断、8个信息子集、旧预报、三种结算和18h时域对照。selected_days为四指定日，figures为PNG/SVG。
''')

def deliver():
    check_protected();dates,a,p,r=inputs();s=dict(np.load(OUT/'experiments/formal_arrays.npz'))
    daily=data_exports(dates,a,p,s)
    metrics,matched,calibration=forecast_diagnostics(dates,a,p,r,s)
    ab,stage=information_value();merged=workbook(dates,p,s)
    figures(dates,a,p,s,daily,metrics,stage);documents(dates,a,p,s,daily,merged,metrics,matched,calibration,ab,stage)
    summary=summarize(s,p);summary.update(pv_curtail_lower_kwh=sum(v['pv_curtail_lower_kwh'] for v in daily[31:]),pv_curtail_upper_kwh=sum(v['pv_curtail_upper_kwh'] for v in daily[31:]))
    write_json(OUT/'summary.json',summary);print('Delivery files generated; independent audit still required.',flush=True)

if __name__=='__main__':deliver()
